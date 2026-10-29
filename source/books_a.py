"""Excel artefacts 04, 05, 07, 09."""
import datetime as dt
from openpyxl.chart import LineChart, Reference, BarChart
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import PROJECT, DOMAINS, FOCUS_AREAS
from xlsx_kit import Book, FILL, font, BORDER, B, HDR_FILL, FX_FILL

CUR = PROJECT["currency"]
MONEY = '#,##0;[Red](#,##0);"-"'
PCT = "0.0%"
DATE = "yyyy-mm-dd"
PRIORITY = ["Critical", "High", "Medium", "Low"]


def summary_block(bk, ws, r, c, title, rows, header=("Item", "Count")):
    """rows: list of (label, formula, fmt)"""
    bk.label(ws, r, c, title, color=B["teal"], size=11).border = None
    for i, h in enumerate(header):
        cell = ws.cell(r + 1, c + i, h); cell.font = font(bold=True, color="FFFFFF"); cell.fill = HDR_FILL; cell.border = BORDER
    for j, row in enumerate(rows):
        lab, f = row[0], row[1]
        fmt = row[2] if len(row) > 2 else None
        bk.label(ws, r + 2 + j, c, lab, bold=False, color=B["ink"])
        bk.value(ws, r + 2 + j, c + 1, f, fmt=fmt, fx=True)
    return r + 2 + len(rows)


# ============================================================ 04 REQUIREMENTS & RTM
def build_04(path):
    bk = Book("04", ["Record every requirement with a unique ID (REQ-###) and a single, testable statement.",
                     "Trace each requirement forward to WBS (doc 05), deliverable and test case; 'Trace complete' turns Yes when all three are filled.",
                     "After baselining, new or changed requirements require a Change Request (doc 20)."])
    cols = [
        dict(h="Req ID", w=10), dict(h="Requirement statement", w=42),
        dict(h="Type", w=16, list=("ReqType", ["Business", "Stakeholder", "Functional", "Non-functional", "Transition", "Project", "Regulatory"])),
        dict(h="Source / stakeholder", w=18), dict(h="Business objective (Charter ref.)", w=16),
        dict(h="Priority (MoSCoW)", w=12, list=("MoSCoW", ["Must", "Should", "Could", "Won't"])),
        dict(h="Acceptance criteria", w=34), dict(h="WBS ID", w=10), dict(h="Deliverable", w=18),
        dict(h="Design / spec ref.", w=14), dict(h="Test case ID", w=12),
        dict(h="Test status", w=12, list=("TestStatus", ["Not run", "Pass", "Fail", "Blocked", "N/A"])),
        dict(h="Owner", w=14),
        dict(h="Requirement status", w=14, list=("ReqStatus", ["Proposed", "Approved", "In build", "Verified", "Accepted", "Deferred", "Rejected"])),
        dict(h="Date raised", w=12, fmt=DATE), dict(h="Comments", w=24),
        dict(h="Trace complete", w=10, f='=IF(A{r}="","",IF(AND(H{r}<>"",I{r}<>"",K{r}<>""),"Yes","No"))', align="center"),
    ]
    ex = [["REQ-001", "The system shall allow authorized users to export the monthly report to PDF.", "Functional",
           "Finance Manager", "O1", "Must", "Export completes < 10 s; PDF matches on-screen layout", "3.2.1",
           "Reporting module", "SPEC-04", "TC-012", "Not run", "[BA]", "Approved", dt.date(2026, 10, 20), ""]]
    ws = bk.register("RTM", cols, ex, n_rows=100, freeze_col=2)
    lr = ws._last_row
    bk.text_format(ws, f"Q5:Q{lr}", {"Yes": (B["green_l"], B["green"]), "No": (B["red_l"], B["red"])})
    bk.text_format(ws, f"L5:L{lr}", {"Pass": (B["green_l"], B["green"]), "Fail": (B["red_l"], B["red"]),
                                      "Blocked": (B["amber_l"], "8A6200")})
    # summary
    s = bk.wb.create_sheet("Summary", 1)
    bk.band(s, "Summary", 6)
    s.column_dimensions["A"].width = 3
    for col, w in zip("BCDEFG", (26, 12, 4, 26, 12, 4)):
        s.column_dimensions[col].width = w
    R = f"RTM!$A$5:$A${lr}"
    r = summary_block(bk, s, 4, 2, "Requirements by status",
                      [(v, f'=COUNTIF(RTM!$N$5:$N${lr},"{v}")') for v in ["Proposed", "Approved", "In build", "Verified", "Accepted", "Deferred", "Rejected"]]
                      + [("Total requirements", f'=COUNTA({R})')])
    summary_block(bk, s, 4, 5, "Requirements by priority",
                  [(v, f'=COUNTIF(RTM!$F$5:$F${lr},"{v}")') for v in ["Must", "Should", "Could", "Won't"]])
    summary_block(bk, s, r + 1, 2, "Traceability & test coverage", [
        ("Fully traced (WBS, deliverable, test)", f'=COUNTIF(RTM!$Q$5:$Q${lr},"Yes")'),
        ("Traceability %", f'=IFERROR(COUNTIF(RTM!$Q$5:$Q${lr},"Yes")/COUNTA({R}),0)', PCT),
        ("Tests passed", f'=COUNTIF(RTM!$L$5:$L${lr},"Pass")'),
        ("Tests failed", f'=COUNTIF(RTM!$L$5:$L${lr},"Fail")'),
        ("Test pass rate (of executed)", f'=IFERROR(COUNTIF(RTM!$L$5:$L${lr},"Pass")/(COUNTIF(RTM!$L$5:$L${lr},"Pass")+COUNTIF(RTM!$L$5:$L${lr},"Fail")),0)', PCT),
    ], header=("Metric", "Value"))
    bk.save(path)


# ============================================================ 05 WBS & DICTIONARY
WBS_SKELETON = [
    ("1", "Project Name (Level 1)", "Project", "Total project scope"),
    ("1.1", "Project Management", "Phase", "All project management effort across the life cycle"),
    ("1.1.1", "Initiation & Charter", "Work Package", "Business case review, charter, stakeholder identification"),
    ("1.1.2", "Planning & Baselining", "Work Package", "Integrated PMP and subsidiary plans, baselines"),
    ("1.1.3", "Monitoring, Control & Reporting", "Work Package", "Status reporting, change control, risk reviews"),
    ("1.1.4", "Project Closure", "Work Package", "Closure report, lessons learned, archive"),
    ("1.2", "Requirements & Analysis", "Phase", "Elicitation, analysis and approval of requirements"),
    ("1.2.1", "Current-State Assessment", "Work Package", ""),
    ("1.2.2", "Requirements Specification", "Work Package", ""),
    ("1.3", "Design", "Phase", "Solution / product design"),
    ("1.3.1", "High-Level Design", "Work Package", ""),
    ("1.3.2", "Detailed Design", "Work Package", ""),
    ("1.4", "Build / Procure", "Phase", "Construction, configuration or procurement of deliverables"),
    ("1.4.1", "[Deliverable A]", "Work Package", ""),
    ("1.4.2", "[Deliverable B]", "Work Package", ""),
    ("1.5", "Test & Quality Control", "Phase", "Verification of deliverables against requirements"),
    ("1.5.1", "System / Integration Testing", "Work Package", ""),
    ("1.5.2", "User Acceptance Testing", "Work Package", ""),
    ("1.6", "Deployment & Transition", "Phase", "Implementation, training and handover to operations"),
    ("1.6.1", "Training & Change Management", "Work Package", ""),
    ("1.6.2", "Go-Live & Hypercare", "Work Package", ""),
    ("1.6.3", "Operational Handover", "Work Package", ""),
]


def build_05(path):
    bk = Book("05", ["The WBS sheet holds the deliverable-oriented hierarchy. Levels and indentation are calculated from the WBS ID (1, 1.1, 1.1.1).",
                     "A generic predictive skeleton is pre-filled — rename, add or remove elements, keeping the 100% rule.",
                     "Complete the WBS Dictionary for every Work Package; it feeds the schedule (doc 07) and budget (doc 09)."])
    etype = ("WBSType", ["Project", "Phase", "Deliverable", "Control Account", "Work Package", "Planning Package"])
    cols = [
        dict(h="WBS ID", w=10),
        dict(h="Level", w=7, f='=IF(A{r}="","",LEN(A{r})-LEN(SUBSTITUTE(A{r},".",""))+1)', align="center"),
        dict(h="Element name", w=34),
        dict(h="WBS (indented view)", w=40, f='=IF(A{r}="","",REPT("      ",B{r}-1)&A{r}&"  "&C{r})', bold=False),
        dict(h="Element type", w=16, list=etype),
        dict(h="Description / scope", w=44), dict(h="Owner", w=16), dict(h="Control account", w=13),
        dict(h="In dictionary?", w=11, f='=IF(OR(A{r}="",E{r}<>"Work Package"),"",IF(COUNTIF(\'WBS Dictionary\'!$A:$A,A{r})>0,"Yes","No"))', align="center"),
    ]
    pre = [[w, None, n, None, t, d] for w, n, t, d in WBS_SKELETON]
    ws = bk.register("WBS", cols, prefill=pre, n_rows=80, freeze_col=1)
    lr = ws._last_row
    ws.conditional_formatting.add(f"A5:H{lr}", FormulaRule(formula=["$B5=1"], fill=FILL(B["navy"]), font=Font(name="Calibri", bold=True, color="FFFFFF")))
    ws.conditional_formatting.add(f"A5:H{lr}", FormulaRule(formula=["$B5=2"], fill=FILL(B["sky"]), font=Font(name="Calibri", bold=True, color=B["navy"])))
    bk.text_format(ws, f"I5:I{lr}", {"No": (B["red_l"], B["red"]), "Yes": (B["green_l"], B["green"])})
    dcols = [dict(h="WBS ID", w=10), dict(h="Work package name", w=26), dict(h="Description of work", w=40),
             dict(h="Deliverable(s)", w=24), dict(h="Acceptance criteria", w=30), dict(h="Assumptions & constraints", w=26),
             dict(h="Responsible (owner)", w=16), dict(h="Milestones", w=18), dict(h="Planned start", w=12, fmt=DATE),
             dict(h="Planned finish", w=12, fmt=DATE), dict(h="Resources", w=18), dict(h="Effort (h)", w=10, fmt="#,##0"),
             dict(h=f"Cost estimate ({CUR})", w=13, fmt=MONEY), dict(h="Quality requirements", w=22),
             dict(h="Technical references", w=18), dict(h="Account code", w=12),
             dict(h="Duration (wd)", w=10, f='=IF(OR(I{r}="",J{r}=""),"",NETWORKDAYS(I{r},J{r}))', align="center")]
    dex = [["1.2.1", "Current-State Assessment", "Interview users and document as-is processes and pain points.",
            "As-is process maps; assessment report", "Report reviewed and accepted by Business Owner", "SMEs available 4 h/week",
            "[BA]", "As-is report approved", dt.date(2026, 11, 16), dt.date(2026, 12, 4), "BA, 3 SMEs", 120, 14400,
            "Peer reviewed; template compliant", "Process standard v2", "CA-1200"]]
    dw = bk.register("WBS Dictionary", dcols, dex, n_rows=60, freeze_col=2)
    bk.save(path)


# ============================================================ 07 MASTER SCHEDULE
def build_07(path):
    bk = Book("07", ["Enter the project start date in cell C3 of the Schedule sheet; the weekly Gantt grid is built from it.",
                     "Enter Start and Duration (working days); Finish is calculated with WORKDAY. Baseline dates are frozen at Gate 2 and only changed via an approved Change Request.",
                     "Finish variance (working days) compares forecast Finish to Baseline Finish: positive = late. Milestones have duration 0."])
    status = ("TaskStatus", ["Not started", "In progress", "Complete", "On hold", "At risk"])
    ttype = ("TaskType", ["Summary", "Task", "Milestone"])
    cols = [
        dict(h="ID", w=5, align="center"), dict(h="WBS ID", w=8), dict(h="Task / milestone name", w=36),
        dict(h="Type", w=10, list=ttype), dict(h="Owner", w=12), dict(h="Pred.", w=7),
        dict(h="Duration (wd)", w=8, fmt="0", align="center"),
        dict(h="Start", w=11, fmt=DATE),
        dict(h="Finish", w=11, fmt=DATE, f='=IF(H{r}="","",IF(N(G{r})=0,H{r},WORKDAY(H{r},G{r}-1)))'),
        dict(h="Baseline start", w=11, fmt=DATE), dict(h="Baseline finish", w=11, fmt=DATE),
        dict(h="Finish var. (wd)", w=8, fmt='0;[Red]-0;0', align="center",
             f='=IF(OR(I{r}="",K{r}=""),"",IF(I{r}>=K{r},NETWORKDAYS(K{r},I{r})-1,-(NETWORKDAYS(I{r},K{r})-1)))'),
        dict(h="% complete", w=8, fmt="0%", dv=("decimal", 0, 1), align="center"),
        dict(h="Status", w=11, list=status), dict(h="Critical path", w=7, list=("YN", ["Y", "N"]), align="center"),
    ]
    start = dt.date.fromisoformat(PROJECT["start"])
    wd = lambda d, n: d + dt.timedelta(days=n)
    pre = [
        [1, "1", PROJECT["name"], "Summary", "[PM]", "", None, start],
        [2, "1.1", "Project Management", "Summary", "[PM]", "", None, start],
        [3, "1.1.1", "Charter approved (Gate 1)", "Milestone", "Sponsor", "", 0, start],
        [4, "1.1.2", "Develop Integrated PMP & baselines", "Task", "[PM]", "3", 20, wd(start, 1)],
        [5, "1.1.2", "PMP baselined (Gate 2)", "Milestone", "Steering Comm.", "4", 0, wd(start, 29)],
        [6, "1.2", "Requirements & Analysis", "Summary", "[BA]", "", None, wd(start, 7)],
        [7, "1.2.1", "Current-state assessment", "Task", "[BA]", "3", 15, wd(start, 7)],
        [8, "1.2.2", "Requirements specification", "Task", "[BA]", "7", 15, wd(start, 28)],
        [9, "1.3", "Design", "Summary", "[Tech Lead]", "", None, wd(start, 49)],
        [10, "1.3.1", "High-level design", "Task", "[Tech Lead]", "8", 10, wd(start, 49)],
        [11, "1.3.2", "Detailed design", "Task", "[Tech Lead]", "10", 15, wd(start, 63)],
        [12, "1.4", "Build / procure deliverables", "Task", "[Team]", "11", 40, wd(start, 84)],
        [13, "1.5", "Testing & UAT", "Task", "[QA Lead]", "12", 20, wd(start, 140)],
        [14, "1.5.2", "Final acceptance (Gate 4)", "Milestone", "Business Owner", "13", 0, wd(start, 168)],
        [15, "1.6", "Deployment & transition", "Task", "[PM]", "14", 15, wd(start, 169)],
        [16, "1.1.4", "Project closure (Gate 5)", "Milestone", "Sponsor", "15", 0, wd(start, 196)],
    ]
    # move any weekend starts to Monday
    for row in pre:
        d = row[7]
        while d.weekday() >= 5:
            d += dt.timedelta(days=1)
        row[7] = d
    ws = bk.register("Schedule", cols, prefill=pre, n_rows=60, freeze_col=3)
    lr = ws._last_row
    # project start input on row 3
    ws.row_dimensions[3].height = 18
    ws["A3"] = "Project start:"; ws["A3"].font = font(bold=True, color=B["navy"], size=9)
    ws.merge_cells("A3:B3")
    ws["C3"] = start; ws["C3"].number_format = DATE; ws["C3"].fill = FILL(B["amber_l"]); ws["C3"].font = font(bold=True)
    ws["D3"] = "Status date:"; ws["D3"].font = font(bold=True, color=B["navy"], size=9)
    ws["E3"] = "=TODAY()"; ws["E3"].number_format = DATE; ws["E3"].font = font(bold=True)
    # summary rows: duration/start/finish rollups left as inputs; baseline = copy of initial plan
    def workday(d, n):
        k = 0
        while k < n:
            d += dt.timedelta(days=1)
            if d.weekday() < 5:
                k += 1
        return d
    fin = {}
    for r, row in zip(range(5, 5 + len(pre)), pre):
        if row[6] is not None:
            fin[r] = row[7] if row[6] == 0 else workday(row[7], row[6] - 1)
    fin[6] = max(fin[7], fin[8], fin[9]); fin[10] = max(fin[11], fin[12]); fin[13] = max(fin[14], fin[15])
    fin[5] = max(v for v in fin.values())
    for r, row in zip(range(5, 5 + len(pre)), pre):
        # baseline = static copy of the initial plan (frozen at Gate 2)
        ws.cell(r, 10).value = row[7]; ws.cell(r, 10).number_format = DATE
        ws.cell(r, 11).value = fin[r]; ws.cell(r, 11).number_format = DATE
        ws.cell(r, 13).value = 0
        ws.cell(r, 14).value = "Not started"
        ws.cell(r, 15).value = "N"
    # summary rows: compute span from children (simple: MIN start / MAX finish of rows below until next summary not handled) -> leave duration blank, finish = manual
    for r, row in zip(range(5, 5 + len(pre)), pre):
        if row[3] == "Summary":
            ws.cell(r, 7).value = None
            ws.cell(r, 9).value = f'=IF(H{r}="","",MAX(I{r + 1}:I{lr}))' if row[1] == "1" else None
            for c in range(1, 16):
                ws.cell(r, c).font = font(bold=True, size=9.5)
    # Phase summaries: finish = max of their own children rows (explicit ranges)
    ws.cell(6, 9).value = "=MAX(I7:I9)"          # Project Management summary (rows 7-9)
    ws.cell(10, 9).value = "=MAX(I11:I12)"       # Requirements
    ws.cell(13, 9).value = "=MAX(I14:I15)"       # Design
    for r in (6, 10, 13):
        ws.cell(r, 9).number_format = DATE
    ws.cell(5, 9).value = f"=MAX(I6:I{lr})"
    # Gantt: 52 weekly columns starting at Monday of project start week
    g0 = 17  # column Q
    weeks = 52
    hdr = 4
    for k in range(weeks):
        c = g0 + k
        L = get_column_letter(c)
        ws.column_dimensions[L].width = 3.2
        cell = ws.cell(hdr, c, f"=$C$3-WEEKDAY($C$3,3)+{7 * k}")
        cell.number_format = "dd-mmm"
        cell.alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")
        cell.font = font(size=7, color="FFFFFF", bold=True); cell.fill = HDR_FILL
        for r in range(1, 3):
            ws.cell(r, c).fill = HDR_FILL
    ws.row_dimensions[hdr].height = 44
    gl = get_column_letter(g0 + weeks - 1)
    G = f"{get_column_letter(g0)}5:{gl}{lr}"
    Q = get_column_letter(g0)
    ws.conditional_formatting.add(G, FormulaRule(formula=[f'AND($D5="Milestone",$H5<>"",{Q}$4<=$H5,{Q}$4+6>=$H5)'], fill=FILL(B["red"]), stopIfTrue=True))
    ws.conditional_formatting.add(G, FormulaRule(formula=[f'AND($D5="Summary",$H5<>"",$I5<>"",{Q}$4<=$I5,{Q}$4+6>=$H5)'], fill=FILL(B["navy"]), stopIfTrue=True))
    ws.conditional_formatting.add(G, FormulaRule(formula=[f'AND($H5<>"",$I5<>"",$O5="Y",{Q}$4<=$I5,{Q}$4+6>=$H5)'], fill=FILL("E67E22"), stopIfTrue=True))
    ws.conditional_formatting.add(G, FormulaRule(formula=[f'AND($H5<>"",$I5<>"",{Q}$4<=$I5,{Q}$4+6>=$H5)'], fill=FILL(B["teal"]), stopIfTrue=True))
    ws.conditional_formatting.add(G, FormulaRule(formula=[f'AND({Q}$4<=$E$3,{Q}$4+6>=$E$3)'], fill=FILL("FFF2CC")))
    for r in range(5, lr + 1):
        for c in range(g0, g0 + weeks):
            ws.cell(r, c).border = BORDER
    ws.auto_filter.ref = f"A4:O{lr}"
    bk.text_format(ws, f"N5:N{lr}", {"Complete": (B["green_l"], B["green"]), "At risk": (B["red_l"], B["red"]),
                                      "In progress": (B["sky"], B["navy"])})
    ws.conditional_formatting.add(f"L5:L{lr}", FormulaRule(formula=["AND(ISNUMBER(L5),L5>10)"], fill=FILL(B["red_l"]), font=Font(name="Calibri", bold=True, color=B["red"])))
    ws.conditional_formatting.add(f"L5:L{lr}", FormulaRule(formula=["AND(ISNUMBER(L5),L5>0)"], fill=FILL(B["amber_l"])))
    # legend
    lg = ws.cell(3, g0, "Legend:"); lg.font = font(size=8, bold=True)
    for i, (txt, col) in enumerate((("Task", B["teal"]), ("Critical", "E67E22"), ("Summary", B["navy"]), ("Milestone", B["red"]), ("This week", "FFF2CC"))):
        c = ws.cell(3, g0 + 2 + i * 4, ""); c.fill = FILL(col)
        t = ws.cell(3, g0 + 3 + i * 4, txt); t.font = font(size=8)

    # Milestone tracker
    m = bk.wb.create_sheet("Milestones", 2)
    bk.band(m, "Milestone Tracker", 8)
    heads = ["Milestone", "Gate", "Baseline", "Forecast / actual", "Variance (wd)", "Status", "Owner", "Comment"]
    widths = [36, 8, 12, 14, 11, 12, 16, 30]
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        m.column_dimensions[get_column_letter(i)].width = w
        cell = m.cell(4, i, h); cell.font = font(bold=True, color="FFFFFF"); cell.fill = HDR_FILL if i not in (3, 4, 5) else FX_FILL
        cell.border = BORDER; cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    mrows = [r for r, row in zip(range(5, 5 + len(pre)), pre) if row[3] == "Milestone"]
    gates = ["G1", "G2", "G4", "G5"]
    for j, sr in enumerate(mrows):
        r = 5 + j
        vals = [f"=Schedule!C{sr}", gates[j] if j < len(gates) else "", f"=Schedule!K{sr}", f"=Schedule!I{sr}",
                f"=Schedule!L{sr}", f"=Schedule!N{sr}", f"=Schedule!E{sr}", ""]
        for i, v in enumerate(vals, start=1):
            c = m.cell(r, i, v); c.border = BORDER; c.font = font(size=9.5)
            if i in (3, 4):
                c.number_format = DATE
            if i in (1, 3, 4, 5, 6, 7):
                c.fill = FILL("EAF2F2")
    m.freeze_panes = "B5"
    bk.save(path)


# ============================================================ 09 PROJECT BUDGET
def build_09(path):
    bk = Book("09", ["Set labour rates and reserve percentages on the Assumptions sheet (amber cells).",
                     "Build the bottom-up estimate on the Cost Estimate sheet by WBS work package; totals roll up to the Summary.",
                     "Enter monthly Planned Value, Earned Value and Actual Cost on Cash Flow & EVM; set the status month to obtain CPI, SPI, EAC and the S-curve."])
    cats = ["Labour", "Contractors / Services", "Hardware / Equipment", "Software / Licences", "Materials",
            "Travel", "Training", "Facilities", "Other"]
    # ---------------- Assumptions
    a = bk.wb.create_sheet("Assumptions")
    bk.band(a, "Assumptions & Rates", 6)
    for col, w in zip("ABCDEF", (30, 16, 14, 4, 30, 14)):
        a.column_dimensions[col].width = w
    bk.label(a, 4, 1, "Parameter", color="FFFFFF", fill=B["navy"]); bk.label(a, 4, 2, "Value", color="FFFFFF", fill=B["navy"])
    params = [("Currency", CUR, None), ("Contingency reserve % (default)", 0.10, PCT),
              ("Management reserve % of cost baseline", 0.05, PCT), ("Hours per working day", 8, "0"),
              ("Estimate class", "Definitive (-5% / +10%)", None)]
    for i, (k, v, f) in enumerate(params):
        bk.label(a, 5 + i, 1, k, bold=False, color=B["ink"])
        c = bk.value(a, 5 + i, 2, v, fmt=f); c.fill = FILL(B["amber_l"])
    bk.label(a, 4, 5, "Role / resource", color="FFFFFF", fill=B["navy"]); bk.label(a, 4, 6, f"Rate / h ({CUR})", color="FFFFFF", fill=B["navy"])
    roles = [("Project Manager", 120), ("Business Analyst", 95), ("Technical Lead", 130), ("Developer / Engineer", 105),
             ("QA / Tester", 85), ("Change / Training Lead", 90), ("Subject-Matter Expert", 80), ("[Role]", 0)]
    for i, (k, v) in enumerate(roles):
        bk.label(a, 5 + i, 5, k, bold=False, color=B["ink"])
        c = bk.value(a, 5 + i, 6, v, fmt=MONEY); c.fill = FILL(B["amber_l"])
    rate_rng = f"Assumptions!$E$5:$F${4 + len(roles)}"

    # ---------------- Cost Estimate
    cols = [
        dict(h="WBS ID", w=9), dict(h="Control account", w=11), dict(h="Cost item / description", w=32),
        dict(h="Category", w=17, list=("CostCat", cats)), dict(h="Role (labour only)", w=18, list=("Roles", [r for r, _ in roles])),
        dict(h="Qty / hours", w=9, fmt="#,##0.0"), dict(h="Unit", w=8),
        dict(h=f"Unit rate ({CUR})", w=11, fmt=MONEY),
        dict(h="Rate used", w=11, fmt=MONEY, f='=IF(H{r}<>"",H{r},IF(E{r}<>"",IFERROR(VLOOKUP(E{r},' + rate_rng + ',2,FALSE),0),0))'),
        dict(h="Base estimate", w=12, fmt=MONEY, f='=IF(F{r}="","",F{r}*I{r})'),
        dict(h="Contingency %", w=9, fmt=PCT, dv=("decimal", 0, 1)),
        dict(h="Contingency", w=11, fmt=MONEY, f='=IF(J{r}="","",J{r}*IF(K{r}="",Assumptions!$B$6,K{r}))'),
        dict(h="Budget (BAC line)", w=12, fmt=MONEY, f='=IF(J{r}="","",J{r}+L{r})', bold=True),
        dict(h="Committed (PO)", w=11, fmt=MONEY), dict(h="Actual to date", w=11, fmt=MONEY),
        dict(h="ETC", w=11, fmt=MONEY),
        dict(h="EAC", w=12, fmt=MONEY, f='=IF(AND(O{r}="",P{r}=""),"",N(O{r})+N(P{r}))'),
        dict(h="Variance (BAC-EAC)", w=12, fmt=MONEY, f='=IF(OR(M{r}="",Q{r}=""),"",M{r}-Q{r})'),
        dict(h="Var. %", w=8, fmt=PCT, f='=IF(OR(R{r}="",N(M{r})=0),"",R{r}/M{r})'),
    ]
    ex = [["1.1.2", "CA-1100", "Project management – planning phase", "Labour", "Project Manager", 160, "h"],
          ["1.2.1", "CA-1200", "Current-state assessment", "Labour", "Business Analyst", 120, "h"],
          ["1.4.1", "CA-1400", "Software licences – year 1", "Software / Licences", None, 25, "user", 600],
          ["1.6.1", "CA-1600", "End-user training delivery", "Training", None, 4, "session", 2500]]
    ws = bk.register("Cost Estimate", cols, ex, n_rows=80, freeze_col=3)
    lr = ws._last_row
    tr = lr + 1
    bk.label(ws, tr, 3, "TOTAL", color="FFFFFF", fill=B["navy"])
    for col in "JLMNOPQR":
        c = ws[f"{col}{tr}"]; c.value = f"=SUM({col}5:{col}{lr})"; c.number_format = MONEY
        c.font = font(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.border = BORDER
    ws.conditional_formatting.add(f"S5:S{lr}", FormulaRule(formula=["AND(ISNUMBER(S5),S5<-0.1)"], fill=FILL(B["red_l"]), font=Font(name="Calibri", bold=True, color=B["red"])))
    ws.conditional_formatting.add(f"S5:S{lr}", FormulaRule(formula=["AND(ISNUMBER(S5),S5<-0.05)"], fill=FILL(B["amber_l"])))

    # ---------------- Summary
    s = bk.wb.create_sheet("Summary", 1)
    bk.band(s, "Budget Summary", 7)
    for col, w in zip("ABCDEFG", (28, 15, 15, 15, 15, 15, 15)):
        s.column_dimensions[col].width = w
    heads = ["Category", "Base estimate", "Contingency", "Budget", "Actual to date", "EAC", "Variance"]
    for i, h in enumerate(heads, start=1):
        c = s.cell(4, i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.border = BORDER
        c.alignment = Alignment(horizontal="center")
    CE = "'Cost Estimate'!"
    for j, cat in enumerate(cats):
        r = 5 + j
        bk.label(s, r, 1, cat, bold=False, color=B["ink"])
        for i, col in enumerate("JLMOQ", start=2):
            bk.value(s, r, i, f"=SUMIF({CE}$D$5:$D${lr},$A{r},{CE}${col}$5:${col}${lr})", fmt=MONEY, fx=True)
        bk.value(s, r, 7, f"=D{r}-F{r}", fmt=MONEY, fx=True)
    t = 5 + len(cats)
    bk.label(s, t, 1, "Cost baseline (incl. contingency)", color="FFFFFF", fill=B["navy"])
    for i in range(2, 8):
        L = get_column_letter(i)
        c = bk.value(s, t, i, f"=SUM({L}5:{L}{t - 1})", fmt=MONEY, fx=True, bold=True)
    bk.label(s, t + 1, 1, "Management reserve", bold=False, color=B["ink"])
    bk.value(s, t + 1, 4, f"=D{t}*Assumptions!$B$7", fmt=MONEY, fx=True)
    bk.label(s, t + 2, 1, "TOTAL PROJECT BUDGET", color="FFFFFF", fill=B["teal"])
    bk.value(s, t + 2, 4, f"=D{t}+D{t + 1}", fmt=MONEY, fx=True, bold=True)
    bk.label(s, t + 3, 1, "Uncategorized lines (check)", bold=False, color=B["grey"])
    bk.value(s, t + 3, 4, f"={CE}M{tr}-D{t}", fmt=MONEY, fx=True)
    ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = "Budget vs EAC by category"
    ch.add_data(Reference(s, min_col=4, min_row=4, max_row=4 + len(cats)), titles_from_data=True)
    ch.add_data(Reference(s, min_col=6, min_row=4, max_row=4 + len(cats)), titles_from_data=True)
    ch.set_categories(Reference(s, min_col=1, min_row=5, max_row=4 + len(cats)))
    ch.height, ch.width = 8, 16
    s.add_chart(ch, f"A{t + 6}")

    # ---------------- Cash Flow & EVM
    c = bk.wb.create_sheet("Cash Flow & EVM", 2)
    months = 12
    bk.band(c, "Time-phased Budget, Earned Value & S-Curve", months + 2)
    c.column_dimensions["A"].width = 30
    for k in range(months + 1):
        c.column_dimensions[get_column_letter(2 + k)].width = 11.5
    start = dt.date.fromisoformat(PROJECT["start"]).replace(day=1)
    bk.label(c, 4, 1, "Month", color="FFFFFF", fill=B["navy"])
    for k in range(months):
        cell = c.cell(4, 2 + k, f"=EDATE(DATE({start.year},{start.month},1),{k})" if k else f"=DATE({start.year},{start.month},1)")
        cell.number_format = "mmm-yy"; cell.font = font(bold=True, color="FFFFFF"); cell.fill = HDR_FILL; cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    tc = 2 + months
    tl = get_column_letter(tc)
    c.cell(4, tc, "Total").font = font(bold=True, color="FFFFFF"); c.cell(4, tc).fill = HDR_FILL
    labels = [("Planned Value (PV) – monthly", True), ("Earned Value (EV) – monthly", True), ("Actual Cost (AC) – monthly", True),
              ("Cumulative PV", False), ("Cumulative EV", False), ("Cumulative AC", False)]
    for j, (lab, inp) in enumerate(labels):
        r = 5 + j
        bk.label(c, r, 1, lab, bold=not inp, color=B["ink"])
        for k in range(months):
            L = get_column_letter(2 + k)
            if inp:
                cell = bk.value(c, r, 2 + k, None, fmt=MONEY)
            else:
                src = r - 3
                cell = bk.value(c, r, 2 + k, f"=SUM($B${src}:{L}{src})" if k else f"=B{src}", fmt=MONEY, fx=True)
        if inp:
            bk.value(c, r, tc, f"=SUM(B{r}:{get_column_letter(tc - 1)}{r})", fmt=MONEY, fx=True, bold=True)
    # check row
    bk.label(c, 11, 1, "PV total vs cost baseline (should be 0)", bold=False, color=B["grey"])
    bk.value(c, 11, tc, f"=Summary!D{t}-{tl}5", fmt=MONEY, fx=True)
    # EVM block
    r0 = 13
    bk.label(c, r0, 1, "EARNED VALUE STATUS", color=B["teal"], size=11).border = None
    bk.label(c, r0 + 1, 1, "Status month (select from row 4)", bold=True, color=B["navy"])
    sm = bk.value(c, r0 + 1, 2, f"=B4", fmt="mmm-yy"); sm.fill = FILL(B["amber_l"])
    mr = f"$B$4:${get_column_letter(tc - 1)}$4"
    pv = f"SUMPRODUCT(({mr}<=$B${r0 + 1})*($B$5:${get_column_letter(tc - 1)}$5))"
    ev = f"SUMPRODUCT(({mr}<=$B${r0 + 1})*($B$6:${get_column_letter(tc - 1)}$6))"
    ac = f"SUMPRODUCT(({mr}<=$B${r0 + 1})*($B$7:${get_column_letter(tc - 1)}$7))"
    evm = [("BAC (cost baseline)", f"=Summary!D{t}", MONEY), ("PV (planned value to date)", "=" + pv, MONEY),
           ("EV (earned value to date)", "=" + ev, MONEY), ("AC (actual cost to date)", "=" + ac, MONEY),
           ("CV = EV − AC", f"=B{r0 + 4}-B{r0 + 5}", MONEY), ("SV = EV − PV", f"=B{r0 + 4}-B{r0 + 3}", MONEY),
           ("CPI = EV / AC", f"=IFERROR(B{r0 + 4}/B{r0 + 5},\"\")", "0.00"), ("SPI = EV / PV", f"=IFERROR(B{r0 + 4}/B{r0 + 3},\"\")", "0.00"),
           ("EAC = BAC / CPI", f"=IFERROR(B{r0 + 2}/B{r0 + 8},B{r0 + 2})", MONEY), ("ETC = EAC − AC", f"=B{r0 + 10}-B{r0 + 5}", MONEY),
           ("VAC = BAC − EAC", f"=B{r0 + 2}-B{r0 + 10}", MONEY),
           ("TCPI = (BAC−EV)/(BAC−AC)", f"=IFERROR((B{r0 + 2}-B{r0 + 4})/(B{r0 + 2}-B{r0 + 5}),\"\")", "0.00"),
           ("Cost status (RAG)", f'=IF(B{r0 + 8}="","",IF(B{r0 + 8}>=0.95,"Green",IF(B{r0 + 8}>=0.9,"Amber","Red")))', None),
           ("Schedule status (RAG)", f'=IF(B{r0 + 9}="","",IF(B{r0 + 9}>=0.95,"Green",IF(B{r0 + 9}>=0.9,"Amber","Red")))', None)]
    for j, (lab, f, fmt) in enumerate(evm):
        bk.label(c, r0 + 2 + j, 1, lab, bold=False, color=B["ink"])
        bk.value(c, r0 + 2 + j, 2, f, fmt=fmt, fx=True, bold=True)
    bk.rag_format(c, f"B{r0 + 14}:B{r0 + 15}")
    lc = LineChart(); lc.title = "S-Curve: cumulative PV, EV, AC"; lc.style = 12
    lc.y_axis.title = CUR; lc.y_axis.numFmt = "#,##0"
    lc.add_data(Reference(c, min_col=1, max_col=1 + months, min_row=8, max_row=10), from_rows=True, titles_from_data=True)
    lc.set_categories(Reference(c, min_col=2, max_col=1 + months, min_row=4))
    lc.x_axis.number_format = "mmm-yy"
    lc.height, lc.width = 9, 20
    c.add_chart(lc, f"D{r0}")
    c.freeze_panes = "B5"
    bk.save(path)
