"""Shared Excel (.xlsx) building blocks: cover sheet, register sheets, lists, styles."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

from config import (PROJECT, DOC_DEFAULTS, BRAND, VERSIONING, TOOLKIT_NAME, TOOLKIT_SUB,
                    STATUSES, CLASSIFICATIONS, doc_id, entry)

B = BRAND
F = B["font"]
thin = Side(style="thin", color=B["line"])
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL = lambda c: PatternFill("solid", start_color=c, end_color=c)
HDR_FILL, FX_FILL = FILL(B["navy"]), FILL(B["teal"])
HEADER_ROW = 4
FIRST = 5


def font(**kw):
    base = dict(name=F, size=10, color=B["ink"])
    base.update(kw)
    return Font(**base)


class Book:
    def __init__(self, no, instructions):
        e = entry(no)
        self.no, self.title, self.code, self.approver = no, e[1], e[2], e[6]
        self.domains, self.focus = e[4], e[5]
        self.id = doc_id(self.code)
        self.wb = Workbook()
        self.lists = {}
        self._cover(instructions)
        self.ls = self.wb.create_sheet("Lists")
        self.ls.sheet_state = "hidden"
        self._list_col = 0

    # ------------------------------------------------------------ cover
    def _cover(self, instructions):
        ws = self.wb.active
        ws.title = "Cover"
        ws.sheet_view.showGridLines = False
        for col, w in zip("ABCDEFG", (2, 24, 18, 18, 30, 16, 14)):
            ws.column_dimensions[col].width = w
        ws.merge_cells("B2:G2"); ws.merge_cells("B3:G3"); ws.merge_cells("B4:G4")
        for r in (2, 3, 4):
            for c in range(2, 8):
                ws.cell(r, c).fill = HDR_FILL
        ws["B2"] = PROJECT["org"].upper(); ws["B2"].font = font(size=9, bold=True, color="9FD3D1")
        ws["B3"] = f"{self.no} — {self.title}"; ws["B3"].font = font(size=20, bold=True, color="FFFFFF")
        ws["B4"] = PROJECT["name"]; ws["B4"].font = font(size=12, color="FFFFFF")
        ws.row_dimensions[3].height = 32
        ws.row_dimensions[2].height = 18; ws.row_dimensions[4].height = 20

        r = 6
        ws.cell(r, 2, "DOCUMENT CONTROL").font = font(bold=True, color=B["teal"], size=11)
        meta = [("Project name", PROJECT["name"]), ("Project ID", PROJECT["id"]), ("Document ID", self.id),
                ("Document title", self.title), ("Version", DOC_DEFAULTS["version"]), ("Status", DOC_DEFAULTS["status"]),
                ("Date", DOC_DEFAULTS["date"]), ("Project Sponsor", PROJECT["sponsor"]),
                ("Project Manager / Owner", PROJECT["pm"]), ("Approver", self.approver),
                ("Classification", PROJECT["classification"]), ("Currency", PROJECT["currency"]),
                ("File location", "[SharePoint / repository path]")]
        for i, (k, v) in enumerate(meta):
            rr = r + 1 + i
            a = ws.cell(rr, 2, k); a.font = font(bold=True, color=B["navy"]); a.fill = FILL(B["sky"]); a.border = BORDER
            ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=4)
            b = ws.cell(rr, 3, v); b.font = font(color=B["grey"] if str(v).startswith("[") else B["ink"])
            b.fill = FILL(B["mist"])
            for c in (3, 4):
                ws.cell(rr, c).border = BORDER
        self.meta_rows = {k: r + 1 + i for i, (k, v) in enumerate(meta)}
        # status validation on cover
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"C{self.meta_rows['Status']}")
        dv2 = DataValidation(type="list", formula1='"' + ",".join(CLASSIFICATIONS) + '"', allow_blank=True)
        ws.add_data_validation(dv2); dv2.add(f"C{self.meta_rows['Classification']}")

        # PMBOK mapping box (right side)
        ws.cell(r, 5, "PMBOK® 8 ALIGNMENT").font = font(bold=True, color=B["teal"], size=11)
        for i, (k, v) in enumerate([("Performance domain(s)", ", ".join(self.domains)),
                                    ("Focus area(s)", ", ".join(self.focus)),
                                    ("Toolkit", TOOLKIT_NAME), ("Standard", TOOLKIT_SUB)]):
            a = ws.cell(r + 1 + i, 5, k); a.font = font(bold=True, color=B["navy"]); a.fill = FILL(B["sky"]); a.border = BORDER
            ws.merge_cells(start_row=r + 1 + i, start_column=6, end_row=r + 1 + i, end_column=7)
            b = ws.cell(r + 1 + i, 6, v); b.font = font(size=9); b.fill = FILL(B["mist"])
            b.alignment = Alignment(wrap_text=True, vertical="center")
            for c in (6, 7):
                ws.cell(r + 1 + i, c).border = BORDER
            ws.row_dimensions[r + 1 + i].height = 28

        # revision history
        r = r + len(meta) + 3
        ws.cell(r, 2, "REVISION HISTORY").font = font(bold=True, color=B["teal"], size=11)
        hdr = ["Version", "Date", "Author", "Description of change", "CR ref.", "Approved by"]
        for i, h in enumerate(hdr):
            c = ws.cell(r + 1, 2 + i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.border = BORDER
        rows = [[DOC_DEFAULTS["version"], DOC_DEFAULTS["date"], DOC_DEFAULTS["author"],
                 "Initial draft created from toolkit template", "N/A", ""]] + [[""] * 6] * 4
        for j, row in enumerate(rows):
            for i, v in enumerate(row):
                c = ws.cell(r + 2 + j, 2 + i, v); c.border = BORDER; c.font = font(size=9)
                if j % 2: c.fill = FILL(B["mist"])
        # versioning legend
        r = r + 2 + len(rows) + 1
        ws.cell(r, 2, "VERSIONING SCHEME").font = font(bold=True, color=B["teal"], size=11)
        for j, (v, m, rule) in enumerate(VERSIONING):
            ws.cell(r + 1 + j, 2, v).font = font(bold=True, size=9)
            ws.cell(r + 1 + j, 3, m).font = font(size=9)
            ws.merge_cells(start_row=r + 1 + j, start_column=4, end_row=r + 1 + j, end_column=7)
            c = ws.cell(r + 1 + j, 4, rule); c.font = font(size=9, color=B["grey"])
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r + 1 + j].height = 26
        # instructions
        r = r + len(VERSIONING) + 2
        ws.cell(r, 2, "HOW TO USE THIS WORKBOOK").font = font(bold=True, color=B["teal"], size=11)
        std = ["Navy column headers = input columns (shaded rows are for data entry). Teal headers marked ƒ = "
               "calculated columns — do not overwrite.",
               "Drop-down lists enforce the toolkit's controlled vocabulary. Rows shown in grey italics are examples — overwrite or delete them.",
               "Update the Version/Status/Date above and add a Revision History row each time the workbook is re-issued."]
        for j, t in enumerate(list(instructions) + std):
            ws.merge_cells(start_row=r + 1 + j, start_column=2, end_row=r + 1 + j, end_column=7)
            c = ws.cell(r + 1 + j, 2, f"{j + 1}. {t}"); c.font = font(size=9.5)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r + 1 + j].height = 30
        ws.sheet_properties.tabColor = B["navy"]
        ws.page_setup.orientation = "portrait"; ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0

    # ------------------------------------------------------------ lists
    def lst(self, name, values):
        """Store a controlled list on the hidden Lists sheet; return absolute range."""
        if name in self.lists:
            return self.lists[name]
        self._list_col += 1
        col = get_column_letter(self._list_col)
        self.ls.cell(1, self._list_col, name).font = font(bold=True)
        for i, v in enumerate(values):
            self.ls.cell(2 + i, self._list_col, v)
        ref = f"Lists!${col}$2:${col}${1 + len(values)}"
        self.lists[name] = ref
        return ref

    # ------------------------------------------------------------ sheet header band
    def band(self, ws, title, ncols, subtitle=None):
        ws.sheet_view.showGridLines = False
        last = get_column_letter(max(ncols, 4))
        ws.merge_cells(f"A1:{last}1"); ws.merge_cells(f"A2:{last}2")
        ws["A1"] = f"{self.no} — {self.title}  |  {title}"
        ws["A1"].font = font(size=14, bold=True, color="FFFFFF")
        ws["A2"] = subtitle or f"{PROJECT['name']}   •   {self.id}   •   v{DOC_DEFAULTS['version']} {DOC_DEFAULTS['status']}   •   {PROJECT['classification']}"
        ws["A2"].font = font(size=9, color="DCE8F2")
        for r in (1, 2):
            for c in range(1, max(ncols, 4) + 1):
                ws.cell(r, c).fill = HDR_FILL
        ws.row_dimensions[1].height = 26
        ws.row_dimensions[2].height = 16
        ws.row_dimensions[3].height = 6

    # ------------------------------------------------------------ register sheet
    def register(self, name, cols, examples=(), n_rows=60, freeze_col=2, tab=None, header_row=HEADER_ROW,
                 subtitle=None, prefill=None):
        """cols: list of dicts {h, w, list?(name, values), f? formula with {r}, fmt?, wrap?, dv?}"""
        ws = self.wb.create_sheet(name)
        self.band(ws, name, len(cols), subtitle)
        hr = header_row
        ws.row_dimensions[hr].height = 34
        for i, c in enumerate(cols, start=1):
            L = get_column_letter(i)
            ws.column_dimensions[L].width = c.get("w", 14)
            is_fx = "f" in c
            cell = ws.cell(hr, i, ("ƒ " if is_fx else "") + c["h"])
            cell.font = font(bold=True, color="FFFFFF", size=9.5)
            cell.fill = FX_FILL if is_fx else HDR_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            cell.border = BORDER
        rows_data = list(prefill or []) + list(examples)
        total = max(n_rows, len(rows_data) + 10)
        for k in range(total):
            r = hr + 1 + k
            data = rows_data[k] if k < len(rows_data) else None
            is_example = data is not None and k >= len(prefill or [])
            for i, c in enumerate(cols, start=1):
                cell = ws.cell(r, i)
                if "f" in c:
                    cell.value = c["f"].format(r=r)
                    cell.fill = FILL("EAF2F2")
                    cell.font = font(size=9.5, bold=c.get("bold", False))
                else:
                    if data is not None and i - 1 < len(data) and data[i - 1] not in (None, ""):
                        cell.value = data[i - 1]
                    cell.font = font(size=9.5, italic=is_example, color=B["grey"] if is_example else B["ink"])
                    if k % 2 == 1:
                        cell.fill = FILL(B["mist"])
                cell.border = BORDER
                cell.alignment = Alignment(wrap_text=c.get("wrap", True), vertical="top",
                                           horizontal=c.get("align", "left" if not c.get("fmt") else "right"))
                if c.get("fmt"):
                    cell.number_format = c["fmt"]
        last_r = hr + total
        # validations
        for i, c in enumerate(cols, start=1):
            L = get_column_letter(i)
            rng = f"{L}{hr + 1}:{L}{last_r}"
            if "list" in c:
                dv = DataValidation(type="list", formula1="=" + self.lst(*c["list"]), allow_blank=True,
                                    showErrorMessage=True, errorTitle="Invalid entry",
                                    error="Please select a value from the list (toolkit controlled vocabulary).")
                ws.add_data_validation(dv); dv.add(rng)
            if "dv" in c:
                kind, lo, hi = c["dv"]
                dv = DataValidation(type=kind, operator="between", formula1=str(lo), formula2=str(hi),
                                    allow_blank=True, showErrorMessage=True,
                                    error=f"Enter a value between {lo} and {hi}.")
                ws.add_data_validation(dv); dv.add(rng)
        ws.freeze_panes = ws.cell(hr + 1, freeze_col + 1)
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{last_r}"
        ws.sheet_properties.tabColor = tab or B["teal"]
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.print_title_rows = f"{hr}:{hr}"
        ws._last_row = last_r
        ws._first_row = hr + 1
        return ws

    # ------------------------------------------------------------ conditional formatting helpers
    def rag_format(self, ws, rng):
        for txt, fill, fc in (("Green", B["green_l"], B["green"]), ("Amber", B["amber_l"], "8A6200"),
                              ("Red", B["red_l"], B["red"])):
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{txt}"'],
                                                          fill=FILL(fill), font=Font(name=F, bold=True, color=fc)))

    def text_format(self, ws, rng, mapping):
        """mapping: {text: (fill, fontcolor)}"""
        for txt, (fill, fc) in mapping.items():
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{txt}"'],
                                                          fill=FILL(fill), font=Font(name=F, bold=True, color=fc)))

    def rating_format(self, ws, rng):
        self.text_format(ws, rng, {"Critical": (B["red"], "FFFFFF"), "High": ("F5B7A8", "7A1E10"),
                                   "Medium": (B["amber_l"], "8A6200"), "Low": (B["green_l"], B["green"])})

    # ------------------------------------------------------------ small helpers
    def label(self, ws, r, c, text, bold=True, color=None, fill=None, size=10):
        cell = ws.cell(r, c, text)
        cell.font = font(bold=bold, color=color or B["navy"], size=size)
        if fill:
            cell.fill = FILL(fill)
        cell.border = BORDER
        return cell

    def value(self, ws, r, c, v, fmt=None, fx=False, bold=False):
        cell = ws.cell(r, c, v)
        cell.font = font(bold=bold)
        cell.fill = FILL("EAF2F2") if fx else FILL(B["mist"])
        cell.border = BORDER
        if fmt:
            cell.number_format = fmt
        return cell

    def save(self, path):
        # move Lists to end
        self.wb.move_sheet("Lists", offset=len(self.wb.sheetnames))
        p = self.wb.properties
        p.title = f"{self.no} - {self.title}"; p.subject = PROJECT["name"]; p.creator = DOC_DEFAULTS["author"]
        p.keywords = f"{self.id}; PMBOK 8"; p.category = TOOLKIT_NAME
        for ws in self.wb.worksheets:
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            if ws.title not in ("Cover",):
                ws.page_setup.orientation = "landscape"
        self.wb.active = 0
        self.wb.save(path)
