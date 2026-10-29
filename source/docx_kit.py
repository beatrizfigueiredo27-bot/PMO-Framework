"""Shared Word (.docx) building blocks: cover, document control, styles, tables."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from config import PROJECT, DOC_DEFAULTS, BRAND, VERSIONING, TOOLKIT_NAME, TOOLKIT_SUB, doc_id, entry

B = BRAND


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ---------------------------------------------------------------- low level
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_margins(table, top=60, bottom=60, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def borders(table, color=None, sz=4):
    color = color or B["line"]
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def field(paragraph, code, placeholder="1"):
    """Insert a Word field (PAGE, NUMPAGES, TOC ...)."""
    run = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = placeholder
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for el in (f1, it, f2, t, f3):
        run._r.append(el)
    return run


def set_cell(cell, text, bold=False, color=None, size=9.5, italic=False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)
        if i < len(lines) - 1:
            r.add_break()
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


# ---------------------------------------------------------------- document
class Doc:
    def __init__(self, no, landscape=False):
        self.no = no
        e = entry(no)
        self.title, self.code, self.approver = e[1], e[2], e[6]
        self.domains, self.focus = e[4], e[5]
        self.id = doc_id(self.code)
        self.d = Document()
        self.landscape = landscape
        self._styles()
        self._page()
        self._header_footer()
        self.sec_no = 0

    # ---- setup
    def _styles(self):
        st = self.d.styles
        n = st["Normal"]
        n.font.name = B["font"]
        n.element.rPr.rFonts.set(qn("w:eastAsia"), B["font"])
        n.font.size = Pt(10.5)
        n.font.color.rgb = rgb(B["ink"])
        n.paragraph_format.space_after = Pt(6)
        n.paragraph_format.line_spacing = 1.12
        for lvl, size, col in ((1, 15, B["navy"]), (2, 12, B["teal"]), (3, 11, B["navy"])):
            h = st[f"Heading {lvl}"]
            h.font.name = B["font_head"]
            rpr = h.element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(a), B["font_head"])
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if fonts.get(qn(a)) is not None:
                    del fonts.attrib[qn(a)]
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = rgb(col)
            h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number"):
            s = st[name]
            s.font.name = B["font"]
            s.font.size = Pt(10.5)
            s.paragraph_format.space_after = Pt(3)

    def _page(self):
        s = self.d.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        if self.landscape:
            s.orientation = WD_ORIENT.LANDSCAPE
            s.page_width, s.page_height = Inches(11), Inches(8.5)
        s.left_margin = s.right_margin = Inches(0.85)
        s.top_margin = Inches(0.9)
        s.bottom_margin = Inches(0.8)
        s.header_distance = Inches(0.35)
        s.footer_distance = Inches(0.35)
        s.different_first_page_header_footer = True

    @property
    def width(self):
        s = self.d.sections[0]
        return s.page_width - s.left_margin - s.right_margin

    def _header_footer(self):
        s = self.d.sections[0]
        # header: project | doc id
        hp = s.header.paragraphs[0]
        hp.text = ""
        tab_pos = self.width
        self._tabs(hp, tab_pos)
        r = hp.add_run(f"{PROJECT['name']}  |  {self.title}")
        r.font.size = Pt(8.5); r.font.color.rgb = rgb(B["grey"])
        r = hp.add_run(f"\t{self.id}  v{DOC_DEFAULTS['version']}")
        r.font.size = Pt(8.5); r.font.color.rgb = rgb(B["navy"]); r.bold = True
        self._bottom_rule(hp, B["teal"])
        # footer: classification | page x of y
        fp = s.footer.paragraphs[0]
        self._tabs(fp, tab_pos)
        r = fp.add_run(f"{PROJECT['org']}  •  Classification: {PROJECT['classification']}  •  {TOOLKIT_NAME}")
        r.font.size = Pt(8); r.font.color.rgb = rgb(B["grey"])
        r = fp.add_run("\tPage "); r.font.size = Pt(8); r.font.color.rgb = rgb(B["grey"])
        rr = field(fp, "PAGE"); rr.font.size = Pt(8)
        r = fp.add_run(" of "); r.font.size = Pt(8); r.font.color.rgb = rgb(B["grey"])
        rr = field(fp, "NUMPAGES"); rr.font.size = Pt(8)

    @staticmethod
    def _tabs(p, pos):
        from docx.enum.text import WD_TAB_ALIGNMENT
        ts = p.paragraph_format.tab_stops
        ts.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CLEAR)
        ts.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.CLEAR)
        ts.add_tab_stop(pos, WD_TAB_ALIGNMENT.RIGHT)

    @staticmethod
    def _bottom_rule(p, color, sz=8):
        pPr = p._p.get_or_add_pPr()
        pb = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz))
        b.set(qn("w:space"), "4"); b.set(qn("w:color"), color)
        pb.append(b); pPr.append(pb)

    # ---- cover + control block
    def cover(self, purpose):
        d = self.d
        for _ in range(3):
            d.add_paragraph()
        p = d.add_paragraph()
        r = p.add_run(PROJECT["org"].upper()); r.font.size = Pt(10); r.bold = True
        r.font.color.rgb = rgb(B["teal"])
        p = d.add_paragraph()
        r = p.add_run(f"DOCUMENT {self.no}"); r.font.size = Pt(10); r.font.color.rgb = rgb(B["grey"])
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(self.title); r.font.size = Pt(30); r.bold = True; r.font.color.rgb = rgb(B["navy"])
        p = d.add_paragraph()
        r = p.add_run(PROJECT["name"]); r.font.size = Pt(16); r.font.color.rgb = rgb(B["teal"])
        self._bottom_rule(p, B["navy"], 18)
        p = d.add_paragraph(); p.paragraph_format.space_before = Pt(10)
        r = p.add_run(purpose); r.font.size = Pt(11); r.italic = True; r.font.color.rgb = rgb(B["grey"])
        d.add_paragraph()
        rows = [
            ("Project name", PROJECT["name"]), ("Project ID", PROJECT["id"]),
            ("Document ID", self.id), ("Version", DOC_DEFAULTS["version"]),
            ("Status", DOC_DEFAULTS["status"]), ("Date", DOC_DEFAULTS["date"]),
            ("Project Sponsor", PROJECT["sponsor"]), ("Project Manager", PROJECT["pm"]),
            ("Approver", self.approver), ("Classification", PROJECT["classification"]),
        ]
        t = d.add_table(rows=0, cols=2)
        borders(t, "FFFFFF", 0); cell_margins(t, 50, 50, 100, 100)
        for k, v in rows:
            c = t.add_row().cells
            set_cell(c[0], k, bold=True, color="FFFFFF", size=9.5); shade(c[0], B["navy"])
            set_cell(c[1], v, size=9.5); shade(c[1], B["mist"])
        self._widths(t, [1.9, 3.4])
        d.add_paragraph()
        p = d.add_paragraph()
        r = p.add_run(f"{TOOLKIT_NAME}  •  {TOOLKIT_SUB}")
        r.font.size = Pt(8.5); r.font.color.rgb = rgb(B["grey"])
        p = d.add_paragraph()
        r = p.add_run("PMBOK 8 alignment — Performance domain(s): " + ", ".join(self.domains)
                      + "  |  Focus area(s): " + ", ".join(self.focus))
        r.font.size = Pt(8.5); r.font.color.rgb = rgb(B["grey"])
        self.page_break()
        self.control_block()

    def control_block(self):
        d = self.d
        self.h("Document Control", numbered=False)
        self.kv_table([
            ("Document ID", self.id), ("Document title", self.title),
            ("Project / ID", f"{PROJECT['name']} / {PROJECT['id']}"),
            ("Current version", DOC_DEFAULTS["version"]), ("Status", DOC_DEFAULTS["status"]),
            ("Owner / Author", DOC_DEFAULTS["author"]), ("Approver", self.approver),
            ("Classification", PROJECT["classification"]),
            ("File location", "[SharePoint / repository path]"),
            ("Next review date", "[yyyy-mm-dd]"),
        ])
        self.h("Revision History", level=2, numbered=False)
        self.table(["Version", "Date", "Author", "Description of change", "CR ref."],
                   [[DOC_DEFAULTS["version"], DOC_DEFAULTS["date"], DOC_DEFAULTS["author"],
                     "Initial draft created from toolkit template", "N/A"], ["", "", "", "", ""], ["", "", "", "", ""]],
                   widths=[0.7, 0.95, 1.4, 2.95, 0.8])
        self.h("Review & Approval", level=2, numbered=False)
        self.table(["Role", "Name", "Signature", "Date", "Decision"],
                   [["Project Sponsor", PROJECT["sponsor"], "", "", "Approve / Reject"],
                    ["Project Manager", PROJECT["pm"], "", "", "Submitted"],
                    ["[Reviewer role]", "", "", "", "Reviewed"]],
                   widths=[1.5, 1.6, 1.4, 0.9, 1.4])
        self.h("Distribution", level=2, numbered=False)
        self.table(["Name / Group", "Role", "Purpose (Action / Information)"],
                   [["Steering Committee", "Governance", "Information"], ["Core project team", "Delivery", "Action"]],
                   widths=[2.4, 2.0, 2.4])
        p = d.add_paragraph()
        r = p.add_run("Versioning: 0.x = draft; 1.0 = first approved baseline; 1.x = minor (no baseline impact); "
                      "2.0+ = major (requires an approved Change Request, doc 20).")
        r.font.size = Pt(8.5); r.italic = True; r.font.color.rgb = rgb(B["grey"])
        self.page_break()
        self.h("Contents", numbered=False)
        p = d.add_paragraph()
        field(p, 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose ‘Update Field’ to build the table of contents.")
        self.page_break()

    # ---- content helpers
    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h(self, text, level=1, numbered=True):
        if level == 1 and numbered:
            self.sec_no += 1
            self.sub_no = 0
            text = f"{self.sec_no}. {text}"
        elif level == 2 and numbered:
            self.sub_no += 1
            text = f"{self.sec_no}.{self.sub_no} {text}"
        return self.d.add_heading(text, level=level)

    def guide(self, text):
        """Blue italic guidance - delete when completing the document."""
        p = self.d.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.3)
        pPr = p._p.get_or_add_pPr()
        pb = OxmlElement("w:pBdr"); b = OxmlElement("w:left")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "12"); b.set(qn("w:space"), "8"); b.set(qn("w:color"), B["teal"])
        pb.append(b); pPr.append(pb)
        r = p.add_run("Guidance: "); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = rgb(B["guide"])
        r = p.add_run(text); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = rgb(B["guide"])
        return p

    def para(self, text, bold=False):
        p = self.d.add_paragraph()
        r = p.add_run(text); r.bold = bold
        return p

    def ph(self, text="[Enter text]"):
        p = self.d.add_paragraph()
        r = p.add_run(text); r.font.color.rgb = rgb(B["grey"])
        return p

    def bullets(self, items, style="List Bullet"):
        for it in items:
            p = self.d.add_paragraph(style=style)
            if isinstance(it, tuple):
                r = p.add_run(it[0] + ": "); r.bold = True
                p.add_run(it[1])
            else:
                p.add_run(it)

    def _widths(self, t, inches):
        t.autofit = False
        total = sum(inches)
        scale = (self.width / 914400) / total if total > (self.width / 914400) else 1
        for row in t.rows:
            for i, w in enumerate(inches):
                if i < len(row.cells):
                    row.cells[i].width = Inches(w * scale)
        # also set grid
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            if i < len(inches):
                gc.set(qn("w:w"), str(int(inches[i] * scale * 1440)))

    def table(self, headers, rows, widths=None, zebra=True, size=9, first_col_bold=False, header_fill=None):
        t = self.d.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(t); cell_margins(t)
        for i, hd in enumerate(headers):
            c = t.rows[0].cells[i]
            set_cell(c, hd, bold=True, color="FFFFFF", size=size)
            shade(c, header_fill or B["navy"])
        repeat_header(t.rows[0])
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                is_ph = isinstance(v, str) and v.startswith("[") and v.endswith("]")
                set_cell(cells[i], v, size=size, bold=(first_col_bold and i == 0),
                         color=B["grey"] if is_ph else None)
                if zebra and ri % 2 == 1:
                    shade(cells[i], B["mist"])
        if widths:
            self._widths(t, widths)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def kv_table(self, pairs, widths=(2.0, 4.8)):
        t = self.d.add_table(rows=0, cols=2)
        borders(t); cell_margins(t)
        for k, v in pairs:
            c = t.add_row().cells
            set_cell(c[0], k, bold=True, size=9.5, color=B["navy"]); shade(c[0], B["sky"])
            is_ph = isinstance(v, str) and v.startswith("[")
            set_cell(c[1], v, size=9.5, color=B["grey"] if is_ph else None)
        self._widths(t, list(widths))
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def approvals(self, extra_roles=None):
        self.h("Approval", numbered=True)
        self.para("By signing below, the approvers confirm they have reviewed this document and accept it as the "
                  "baseline for the project. Subsequent baseline changes follow integrated change control.")
        roles = [("Project Sponsor", PROJECT["sponsor"]), ("Project Manager", PROJECT["pm"])] + (extra_roles or [])
        self.table(["Role", "Name", "Signature", "Date"], [[r, n, "", ""] for r, n in roles],
                   widths=[1.8, 2.0, 2.0, 1.0])

    def save(self, path):
        # ask Word to refresh TOC/fields on open
        settings = self.d.settings.element
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
        settings.append(uf)
        self.d.core_properties.title = f"{self.no} - {self.title}"
        self.d.core_properties.subject = PROJECT["name"]
        self.d.core_properties.author = DOC_DEFAULTS["author"]
        self.d.core_properties.keywords = f"{self.id}; PMBOK 8; {', '.join(self.domains)}"
        self.d.core_properties.category = TOOLKIT_NAME
        self.d.save(path)
