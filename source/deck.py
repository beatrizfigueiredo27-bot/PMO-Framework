"""22 - Project Status Report (.pptx)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from config import PROJECT, DOC_DEFAULTS, BRAND, TOOLKIT_NAME, TOOLKIT_SUB, doc_id, entry, VERSIONING

B = BRAND
C = lambda h: RGBColor.from_string(h)
W, H = Inches(13.333), Inches(7.5)
ID = doc_id("PSR")


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = C(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = C(line); s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, t, size=12, bold=False, color=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = B["font"]
        r.font.color.rgb = C(color or B["ink"])
    return tb


def table(slide, x, y, w, headers, rows, col_w=None, size=10, row_h=0.36):
    n = len(rows) + 1
    shp = slide.shapes.add_table(n, len(headers), x, y, w, Inches(row_h * n))
    t = shp.table
    tot = sum(col_w) if col_w else len(headers)
    for i in range(len(headers)):
        t.columns[i].width = int(w * ((col_w[i] if col_w else 1) / tot))
    for i, h in enumerate(headers):
        c = t.cell(0, i); c.text = h
        c.fill.solid(); c.fill.fore_color.rgb = C(B["navy"])
        p = c.text_frame.paragraphs[0]; p.runs[0].font.size = Pt(size); p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = C("FFFFFF"); p.runs[0].font.name = B["font"]
    for r, row in enumerate(rows, start=1):
        for i, v in enumerate(row):
            c = t.cell(r, i); c.text = str(v)
            c.fill.solid(); c.fill.fore_color.rgb = C(B["mist"] if r % 2 == 0 else "FFFFFF")
            p = c.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].font.size = Pt(size); p.runs[0].font.name = B["font"]
                p.runs[0].font.color.rgb = C(B["grey"] if str(v).startswith("[") else B["ink"])
            ragc = {"Green": (B["green_l"], B["green"]), "Amber": (B["amber_l"], "8A6200"), "Red": (B["red_l"], B["red"])}
            if v in ragc:
                c.fill.fore_color.rgb = C(ragc[v][0]); p.runs[0].font.color.rgb = C(ragc[v][1]); p.runs[0].font.bold = True
    return t


class Deck:
    def __init__(self):
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = W, H
        self.blank = self.p.slide_layouts[6]
        self.n = 0

    def slide(self, title, subtitle=None):
        s = self.p.slides.add_slide(self.blank)
        self.n += 1
        box(s, 0, 0, W, Inches(0.95), B["navy"])
        box(s, 0, Inches(0.95), W, Inches(0.06), B["teal"])
        text(s, Inches(0.5), Inches(0.17), Inches(9.5), Inches(0.6), title, 24, True, "FFFFFF", anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(9.3), Inches(0.2), Inches(3.6), Inches(0.6),
             [PROJECT["name"], "Status date: [yyyy-mm-dd]"], 10, False, "DCE8F2", PP_ALIGN.RIGHT)
        if subtitle:
            text(s, Inches(0.5), Inches(1.1), Inches(12.3), Inches(0.4), subtitle, 11, False, B["grey"], italic=True)
        # footer
        box(s, Inches(0.5), Inches(7.05), Inches(12.33), Emu(9525), B["line"])
        text(s, Inches(0.5), Inches(7.1), Inches(9), Inches(0.3),
             f"{PROJECT['org']}  •  {ID}  v{DOC_DEFAULTS['version']}  •  Classification: {PROJECT['classification']}", 8, color=B["grey"])
        text(s, Inches(11.3), Inches(7.1), Inches(1.53), Inches(0.3), f"{self.n}", 8, True, B["navy"], PP_ALIGN.RIGHT)
        return s

    def rag_tile(self, s, x, y, w, h, label, status, note=""):
        cols = {"Green": B["green"], "Amber": B["amber"], "Red": B["red"]}
        box(s, x, y, w, h, "FFFFFF", B["line"])
        box(s, x, y, Inches(0.12), h, cols[status])
        text(s, x + Inches(0.2), y + Inches(0.08), w - Inches(0.3), Inches(0.3), label.upper(), 9, True, B["grey"])
        text(s, x + Inches(0.2), y + Inches(0.35), w - Inches(0.3), Inches(0.4), status, 18, True, cols[status])
        if note:
            text(s, x + Inches(0.2), y + Inches(0.8), w - Inches(0.3), h - Inches(0.85), note, 9, color=B["grey"])


def build_22_pptx(path):
    d = Deck()
    # 1 title
    s = d.p.slides.add_slide(d.blank); d.n += 1
    box(s, 0, 0, W, H, B["navy"])
    box(s, 0, Inches(5.1), W, Inches(0.08), B["teal"])
    text(s, Inches(0.8), Inches(1.0), Inches(10), Inches(0.4), PROJECT["org"].upper(), 12, True, "9FD3D1")
    text(s, Inches(0.8), Inches(1.5), Inches(11), Inches(1.0), "Project Status Report", 44, True, "FFFFFF")
    text(s, Inches(0.8), Inches(2.55), Inches(11), Inches(0.6), PROJECT["name"], 24, False, "DCE8F2")
    text(s, Inches(0.8), Inches(3.4), Inches(11), Inches(0.5), "Reporting period: [yyyy-mm-dd] to [yyyy-mm-dd]   •   Report #[nn]", 14, False, "FFFFFF")
    meta = [("Project ID", PROJECT["id"]), ("Document ID", ID), ("Version", DOC_DEFAULTS["version"]),
            ("Status", DOC_DEFAULTS["status"]), ("Sponsor", PROJECT["sponsor"]), ("Project Manager", PROJECT["pm"]),
            ("Classification", PROJECT["classification"])]
    for i, (k, v) in enumerate(meta):
        x = Inches(0.8 + (i % 4) * 3.0); y = Inches(5.45 + (i // 4) * 0.7)
        text(s, x, y, Inches(2.9), Inches(0.25), k.upper(), 9, True, "9FD3D1")
        text(s, x, y + Inches(0.24), Inches(2.9), Inches(0.35), v, 12, False, "FFFFFF")
    text(s, Inches(0.8), Inches(6.9), Inches(11), Inches(0.3), f"{TOOLKIT_NAME}  •  {TOOLKIT_SUB}", 9, False, "9FB3C8")

    # 2 executive summary
    s = d.slide("Executive Summary", "Overall status against baselines using the RAG thresholds defined in the Integrated PMP (doc 02).")
    tiles = [("Overall", "Green", "[One-line rationale]"), ("Scope", "Green", "No unapproved change"), ("Schedule", "Amber", "SPI 0.93; milestone M3 +6 wd"),
             ("Cost", "Green", "CPI 1.02"), ("Risk", "Amber", "1 critical risk with plan"), ("Quality", "Green", "All criteria met"),
             ("Resources", "Green", "Fully staffed"), ("Stakeholders", "Green", "Engagement on target")]
    for i, (lab, st, note) in enumerate(tiles):
        x = Inches(0.5 + (i % 4) * 3.1); y = Inches(1.6 + (i // 4) * 1.35)
        d.rag_tile(s, x, y, Inches(2.95), Inches(1.2), lab, st, note)
    box(s, Inches(0.5), Inches(4.4), Inches(12.33), Inches(2.5), B["mist"])
    text(s, Inches(0.7), Inches(4.5), Inches(6), Inches(0.35), "Summary", 13, True, B["navy"])
    text(s, Inches(0.7), Inches(4.9), Inches(11.9), Inches(1.9),
         ["[Progress this period and key achievements]", "[Main concerns and corrective actions]",
          "[Decisions or support required from the Sponsor / Steering Committee]"], 12, color=B["grey"])

    # 3 milestones
    s = d.slide("Milestones & Schedule", "Baseline vs forecast; variance in working days (positive = late). Source: Master Schedule (doc 07).")
    table(s, Inches(0.5), Inches(1.6), Inches(12.33), ["Milestone", "Gate", "Baseline", "Forecast / actual", "Var. (wd)", "Status", "Comment"],
          [["Charter approved", "G1", "[date]", "[date]", "0", "Green", "Complete"],
           ["PMP baselined", "G2", "[date]", "[date]", "0", "Green", ""],
           ["[Design complete]", "", "[date]", "[date]", "+6", "Amber", "[Reason / recovery]"],
           ["Final acceptance", "G4", "[date]", "[date]", "0", "Green", ""],
           ["Project closure", "G5", "[date]", "[date]", "0", "Green", ""]],
          col_w=[3, 0.7, 1.3, 1.6, 1, 1, 3.4], size=11, row_h=0.45)
    # mini KPI
    for i, (k, v) in enumerate((("SPI", "[0.93]"), ("Critical path float", "[0 d]"), ("% complete (planned / actual)", "[45% / 42%]"))):
        x = Inches(0.5 + i * 4.15)
        box(s, x, Inches(5.0), Inches(3.95), Inches(1.6), "FFFFFF", B["line"])
        text(s, x + Inches(0.2), Inches(5.1), Inches(3.6), Inches(0.3), k.upper(), 10, True, B["grey"])
        text(s, x + Inches(0.2), Inches(5.5), Inches(3.6), Inches(0.8), v, 30, True, B["navy"])

    # 4 financial
    s = d.slide("Financial & Earned Value Performance", f"All figures in {PROJECT['currency']}. Source: Project Budget (doc 09) — Cash Flow & EVM sheet.")
    kpis = [("BAC", "[0]"), ("EAC", "[0]"), ("CPI", "[1.02]"), ("SPI", "[0.93]"), ("VAC", "[0]"), ("Reserve used", "[0%]")]
    for i, (k, v) in enumerate(kpis):
        x = Inches(0.5 + i * 2.07)
        box(s, x, Inches(1.6), Inches(1.95), Inches(1.3), B["navy"] if i < 2 else "FFFFFF", None if i < 2 else B["line"])
        text(s, x + Inches(0.15), Inches(1.7), Inches(1.7), Inches(0.3), k, 10, True, "9FD3D1" if i < 2 else B["grey"])
        text(s, x + Inches(0.15), Inches(2.05), Inches(1.7), Inches(0.6), v, 24, True, "FFFFFF" if i < 2 else B["navy"])
    table(s, Inches(0.5), Inches(3.2), Inches(12.33), ["Metric", "To date", "Previous period", "Threshold", "Status", "Commentary"],
          [["Planned Value (PV)", "[ ]", "[ ]", "", "", ""], ["Earned Value (EV)", "[ ]", "[ ]", "", "", ""],
           ["Actual Cost (AC)", "[ ]", "[ ]", "", "", ""], ["Cost Performance Index (CPI)", "[ ]", "[ ]", "≥ 0.95", "Green", ""],
           ["Schedule Performance Index (SPI)", "[ ]", "[ ]", "≥ 0.95", "Amber", ""], ["Estimate at Completion (EAC)", "[ ]", "[ ]", "≤ BAC", "Green", ""]],
          col_w=[3, 1.4, 1.6, 1.2, 1, 4], size=10.5, row_h=0.42)

    # 5 risks & issues
    s = d.slide("Top Risks & Issues", "Open risks scored ≥ 10 (High/Critical) and open High/Critical issues. Sources: docs 16 and 19.")
    table(s, Inches(0.5), Inches(1.6), Inches(12.33), ["ID", "Type", "Description", "P×I / priority", "Rating", "Response / action", "Owner", "Due"],
          [["R-001", "Risk", "[Supplier delivers key component late]", "3×4 = 12", "High", "[Weekly check-ins; alternate supplier]", "[PM]", "[date]"],
           ["R-00x", "Risk", "", "", "", "", "", ""], ["I-001", "Issue", "[Test environment unavailable]", "High", "", "[Escalated to IT ops]", "[Tech Lead]", "[date]"],
           ["I-00x", "Issue", "", "", "", "", "", ""]],
          col_w=[0.8, 0.8, 3.4, 1.3, 1, 3.2, 1.1, 1], size=10.5, row_h=0.55)
    text(s, Inches(0.5), Inches(5.0), Inches(12), Inches(0.3), "Risk profile (open risks)", 12, True, B["navy"])
    for i, (lab, col) in enumerate((("Critical", B["red"]), ("High", "E67E22"), ("Medium", B["amber"]), ("Low", B["green"]))):
        x = Inches(0.5 + i * 3.1)
        box(s, x, Inches(5.4), Inches(2.95), Inches(1.2), col)
        text(s, x + Inches(0.2), Inches(5.45), Inches(2.6), Inches(0.3), lab.upper(), 10, True, "FFFFFF")
        text(s, x + Inches(0.2), Inches(5.75), Inches(2.6), Inches(0.7), "[n]", 28, True, "FFFFFF")

    # 6 changes & decisions
    s = d.slide("Changes & Decisions Required", "Change requests from the Change Log (doc 21); decisions recorded in the Decision Log (doc 23).")
    text(s, Inches(0.5), Inches(1.55), Inches(6), Inches(0.35), "Change requests", 13, True, B["navy"])
    table(s, Inches(0.5), Inches(1.95), Inches(12.33), ["CR ID", "Title", "Schedule (wd)", f"Cost ({PROJECT['currency']})", "Authority", "Status"],
          [["CR-001", "[Add mobile access to reporting module]", "+8", "18,000", "CCB", "Pending decision"], ["CR-00x", "", "", "", "", ""]],
          col_w=[1, 4.5, 1.3, 1.4, 1.5, 2], size=10.5, row_h=0.45)
    text(s, Inches(0.5), Inches(3.75), Inches(6), Inches(0.35), "Decisions required from Sponsor / Steering Committee", 13, True, B["navy"])
    table(s, Inches(0.5), Inches(4.15), Inches(12.33), ["#", "Decision required", "Options", "Recommendation", "Needed by"],
          [["1", "[Decision]", "[A / B]", "[Recommended option]", "[date]"], ["2", "", "", "", ""]],
          col_w=[0.5, 4, 3, 3, 1.3], size=10.5, row_h=0.5)

    # 7 accomplishments
    s = d.slide("Accomplishments & Next Period")
    for i, (hd, col) in enumerate((("Accomplished this period", B["teal"]), ("Planned for next period", B["navy"]))):
        x = Inches(0.5 + i * 6.25)
        box(s, x, Inches(1.4), Inches(6.08), Inches(0.5), col)
        text(s, x + Inches(0.2), Inches(1.45), Inches(5.6), Inches(0.4), hd, 14, True, "FFFFFF", anchor=MSO_ANCHOR.MIDDLE)
        box(s, x, Inches(1.9), Inches(6.08), Inches(4.9), B["mist"])
        text(s, x + Inches(0.25), Inches(2.1), Inches(5.6), Inches(4.5),
             ["•  [Deliverable / milestone completed]", "•  [Key activity completed]", "•  [Stakeholder event held]", "•  [ ]"]
             if i == 0 else ["•  [Planned deliverable / activity]", "•  [Upcoming milestone]", "•  [Upcoming decision / gate]", "•  [ ]"],
             13, color=B["grey"])

    # 8 appendix doc control
    s = d.slide("Appendix — Document Control & RAG Definitions")
    table(s, Inches(0.5), Inches(1.4), Inches(6.0), ["Field", "Value"],
          [["Document ID", ID], ["Project / ID", f"{PROJECT['name']} / {PROJECT['id']}"], ["Version / status", f"{DOC_DEFAULTS['version']} / {DOC_DEFAULTS['status']}"],
           ["Owner", PROJECT["pm"]], ["Approver", entry("22")[6]], ["Classification", PROJECT["classification"]], ["Distribution", "Per Communication Plan (doc 12)"]],
          col_w=[1.6, 4.4], size=10, row_h=0.4)
    table(s, Inches(6.83), Inches(1.4), Inches(6.0), ["Version", "Date", "Author", "Change"],
          [[DOC_DEFAULTS["version"], DOC_DEFAULTS["date"], DOC_DEFAULTS["author"], "Initial template"], ["", "", "", ""], ["", "", "", ""]],
          col_w=[0.9, 1.2, 1.6, 2.3], size=10, row_h=0.4)
    table(s, Inches(6.83), Inches(3.3), Inches(6.0), ["RAG", "Definition"],
          [["Green", "On track; within tolerance"], ["Amber", "At risk; tolerance threatened; PM corrective action"],
           ["Red", "Tolerance breached; escalation to Sponsor / Steering Committee"]], col_w=[1, 5], size=10, row_h=0.42)
    d.p.core_properties.title = "22 - Project Status Report"; d.p.core_properties.subject = PROJECT["name"]
    d.p.core_properties.author = DOC_DEFAULTS["author"]; d.p.core_properties.keywords = f"{ID}; PMBOK 8"
    d.p.save(path)
