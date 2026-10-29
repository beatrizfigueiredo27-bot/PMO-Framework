"""Excel artefacts 10, 12, 14, 16, 19, 21, 23, 24."""
import datetime as dt
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from config import PROJECT, DOMAINS, FOCUS_AREAS
from xlsx_kit import Book, FILL, font, BORDER, B, HDR_FILL, FX_FILL
from books_a import summary_block, MONEY, PCT, DATE, PRIORITY

CUR = PROJECT["currency"]
D = dt.date(2026, 10, 20)
ENG = ["Unaware", "Resistant", "Neutral", "Supportive", "Leading"]
OPEN_STATUS = ("ItemStatus", ["Open", "In progress", "On hold", "Closed"])
SCORE = ("Score5", [1, 2, 3, 4, 5])


def rating_f(col):
    return f'=IF({col}{{r}}="","",IF({col}{{r}}>=16,"Critical",IF({col}{{r}}>=10,"High",IF({col}{{r}}>=5,"Medium","Low"))))'


def status_fmt(bk, ws, rng):
    bk.text_format(ws, rng, {"Closed": ("E5E8EB", B["grey"]), "Open": (B["sky"], B["navy"]),
                             "On hold": (B["amber_l"], "8A6200")})


# ============================================================ 10 STAKEHOLDER REGISTER
def build_10(path):
    bk = Book("10", ["Identify every individual or group that can affect or be affected by the project.",
                     "Score power/influence and interest from 1 (low) to 5 (high); the quadrant and engagement gap are calculated.",
                     "Engagement strategies and actions are detailed in the Stakeholder Engagement Plan (doc 11); communication logistics in doc 12."])
    cols = [dict(h="ID", w=7), dict(h="Name", w=20), dict(h="Title / role", w=20), dict(h="Organization / unit", w=18),
            dict(h="Internal / external", w=10, list=("IntExt", ["Internal", "External"])), dict(h="Contact (email / phone)", w=22),
            dict(h="Key requirements / expectations", w=32),
            dict(h="Power / influence (1-5)", w=9, list=SCORE, align="center"), dict(h="Interest (1-5)", w=9, list=SCORE, align="center"),
            dict(h="Quadrant", w=15, f='=IF(OR(H{r}="",I{r}=""),"",IF(H{r}>=3,IF(I{r}>=3,"Manage closely","Keep satisfied"),IF(I{r}>=3,"Keep informed","Monitor")))'),
            dict(h="Attitude", w=11, list=("Attitude", ["Champion", "Supporter", "Neutral", "Critic", "Blocker"])),
            dict(h="Current engagement", w=12, list=("Eng", ENG)), dict(h="Desired engagement", w=12, list=("Eng", ENG)),
            dict(h="Engagement gap", w=10, f='=IF(OR(L{r}="",M{r}=""),"",IF(MATCH(L{r},' + "{rng}" + ',0)<MATCH(M{r},' + "{rng}" + ',0),"Gap","On target"))'),
            dict(h="Engagement strategy / key actions", w=32), dict(h="Relationship owner", w=15),
            dict(h="Comms frequency", w=12, list=("Freq", ["Daily", "Weekly", "Bi-weekly", "Monthly", "Quarterly", "As needed"])),
            dict(h="Last updated", w=11, fmt=DATE)]
    rng = bk.lst("Eng", ENG)
    cols[13]["f"] = cols[13]["f"].replace("{rng}", rng.replace("{", "{{"))
    ex = [["SH-01", PROJECT["sponsor"], "Project Sponsor", "[Executive office]", "Internal", "", "ROI realized; delivery on time",
           5, 5, None, "Champion", "Supportive", "Leading", None, "Bi-weekly 1:1; chair Steering Committee", PROJECT["pm"], "Bi-weekly", D],
          ["SH-02", "[End-user group]", "Operational users", "[Operations]", "Internal", "", "Minimal disruption; usable tool",
           2, 4, None, "Neutral", "Unaware", "Supportive", None, "Demos, change champions, training", "[Change lead]", "Monthly", D]]
    ws = bk.register("Register", cols, ex, n_rows=60, freeze_col=2)
    lr = ws._last_row
    bk.text_format(ws, f"N5:N{lr}", {"Gap": (B["amber_l"], "8A6200"), "On target": (B["green_l"], B["green"])})
    bk.text_format(ws, f"J5:J{lr}", {"Manage closely": (B["red_l"], B["red"]), "Keep satisfied": (B["amber_l"], "8A6200"),
                                      "Keep informed": (B["sky"], B["navy"]), "Monitor": ("E5E8EB", B["grey"])})
    # power-interest grid
    g = bk.wb.create_sheet("Power-Interest Grid", 2)
    bk.band(g, "Power / Interest Grid", 6)
    g.column_dimensions["A"].width = 4; g.column_dimensions["B"].width = 14
    for col in "CD":
        g.column_dimensions[col].width = 38
    g.merge_cells("B5:B6"); g["B5"] = "HIGH\nPOWER"; g["B7"] = "LOW\nPOWER"; g.merge_cells("B7:B8")
    for c in ("B5", "B7"):
        g[c].font = font(bold=True, color=B["navy"]); g[c].alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    g["C9"] = "LOW INTEREST"; g["D9"] = "HIGH INTEREST"
    for c in ("C9", "D9"):
        g[c].font = font(bold=True, color=B["navy"]); g[c].alignment = Alignment(horizontal="center")
    quads = {("C", 5): ("KEEP SATISFIED", "Keep satisfied", B["amber_l"]), ("D", 5): ("MANAGE CLOSELY", "Manage closely", B["red_l"]),
             ("C", 7): ("MONITOR", "Monitor", "E5E8EB"), ("D", 7): ("KEEP INFORMED", "Keep informed", B["sky"])}
    for (col, r), (title, key, fill) in quads.items():
        g[f"{col}{r}"] = title; g[f"{col}{r}"].font = font(bold=True, size=12, color=B["navy"])
        g[f"{col}{r + 1}"] = f'=COUNTIF(Register!$J$5:$J${lr},"{key}")&" stakeholder(s)"'
        g[f"{col}{r + 1}"].font = font(size=20, bold=True, color=B["teal"])
        for rr in (r, r + 1):
            g[f"{col}{rr}"].fill = FILL(fill); g[f"{col}{rr}"].border = BORDER
            g[f"{col}{rr}"].alignment = Alignment(horizontal="center", vertical="center")
        g.row_dimensions[r].height = 30; g.row_dimensions[r + 1].height = 60
    g["C11"] = "Engagement gaps to close:"; g["C11"].font = font(bold=True)
    g["D11"] = f'=COUNTIF(Register!$N$5:$N${lr},"Gap")'; g["D11"].font = font(bold=True, size=14, color=B["red"])
    bk.save(path)


# ============================================================ 12 COMMUNICATION PLAN
def build_12(path):
    bk = Book("12", ["One row per communication/report/meeting. Typical predictive-project communications are pre-filled — tailor them.",
                     "Audience should reference stakeholder IDs or groups from the Stakeholder Register (doc 10).",
                     "Escalation paths and RAG thresholds come from the Integrated PMP (doc 02)."])
    meth = ("Method", ["Meeting (in person)", "Meeting (virtual)", "Email", "Report", "Dashboard", "Intranet / SharePoint", "Workshop", "Newsletter", "Presentation"])
    freq = ("Freq", ["Once", "Daily", "Weekly", "Bi-weekly", "Monthly", "Quarterly", "Per gate", "As needed"])
    cols = [dict(h="ID", w=7), dict(h="Communication / artefact", w=24), dict(h="Purpose / objective", w=30),
            dict(h="Audience", w=22), dict(h="Key content", w=30), dict(h="Method / channel", w=16, list=meth),
            dict(h="Frequency", w=11, list=freq), dict(h="Timing (day / time)", w=14), dict(h="Owner / sender", w=14),
            dict(h="Format / template", w=18), dict(h="Feedback / approval", w=18), dict(h="Storage location", w=18),
            dict(h="Status", w=10, list=("CStatus", ["Planned", "Active", "Retired"]))]
    pre = [
        ["C-01", "Project kick-off meeting", "Align team & stakeholders on objectives, scope, roles", "Sponsor, team, key stakeholders", "Charter, approach, roles, next steps", "Meeting (in person)", "Once", "Week 1", PROJECT["pm"], "Kick-off deck", "Attendance & Q&A", "[Repository]", "Planned"],
        ["C-02", "Team stand-up / coordination", "Coordinate daily work, surface blockers", "Core team", "Progress, plans, blockers", "Meeting (virtual)", "Daily", "09:00, 15 min", PROJECT["pm"], "Action list", "Verbal", "", "Planned"],
        ["C-03", "Weekly team status meeting", "Review progress, risks, issues, actions", "Core team, workstream leads", "Schedule, RAID, actions", "Meeting (virtual)", "Weekly", "Tuesday 10:00", PROJECT["pm"], "Agenda & minutes", "Minutes circulated in 24 h", "[Repository]", "Planned"],
        ["C-04", "Project Status Report", "Report performance vs baselines", "Sponsor, Steering Committee, PMO", "RAG, milestones, EVM, top risks, decisions", "Report", "Bi-weekly", "Friday", PROJECT["pm"], "Doc 22", "Sponsor comments", "[Repository]", "Planned"],
        ["C-05", "Steering Committee meeting", "Governance, decisions, gate approvals", "Steering Committee", "Status, escalations, CRs, gate decisions", "Meeting (in person)", "Monthly", "[1st Thursday]", PROJECT["sponsor"], "SC pack", "Decisions logged (doc 23)", "[Repository]", "Planned"],
        ["C-06", "Change Control Board", "Decide on change requests", "CCB members", "CR impact analyses", "Meeting (virtual)", "As needed", "", PROJECT["pm"], "Doc 20", "CCB decision", "[Repository]", "Planned"],
        ["C-07", "Sponsor 1:1", "Escalations, steering, relationship", "Sponsor", "Key issues, decisions needed", "Meeting (virtual)", "Bi-weekly", "", PROJECT["pm"], "Informal", "Verbal", "", "Planned"],
        ["C-08", "Stakeholder newsletter", "Build awareness and support", "Wider organization, end users", "Progress, what's changing, when", "Newsletter", "Monthly", "", "[Change lead]", "Newsletter template", "Survey / feedback form", "[Intranet]", "Planned"],
        ["C-09", "Gate review", "Formal phase approval", "Steering Committee", "Gate checklist, deliverables, readiness", "Presentation", "Per gate", "", PROJECT["pm"], "Gate checklist", "Signed approval", "[Repository]", "Planned"],
        ["C-10", "Lessons learned session", "Capture and share lessons", "Team, key stakeholders", "What worked / to improve", "Workshop", "Per gate", "", PROJECT["pm"], "Doc 24", "Register updated", "[Repository]", "Planned"],
        ["C-11", "Project closure communication", "Announce completion & handover", "All stakeholders", "Outcomes, benefits, thanks, contacts", "Email", "Once", "At closure", PROJECT["sponsor"], "Doc 25 summary", "", "[Repository]", "Planned"],
    ]
    ws = bk.register("Communication Plan", cols, prefill=pre, n_rows=40, freeze_col=2)
    # escalation sheet
    e = bk.wb.create_sheet("Escalation Path", 2)
    bk.band(e, "Escalation Path & Meeting Cadence", 5)
    heads = ["Level", "Escalate to", "Trigger", "Channel", "Expected response"]
    for i, (h, w) in enumerate(zip(heads, (8, 22, 44, 20, 18)), start=1):
        e.column_dimensions[get_column_letter(i)].width = w
        c = e.cell(4, i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.border = BORDER
    rows = [["1", "Project Manager", "Issue resolvable within team tolerance", "Stand-up / direct", "2 business days"],
            ["2", "Project Sponsor", "Threatens tolerance, cross-functional or resource conflict", "Sponsor 1:1 / email", "5 business days"],
            ["3", "Steering Committee", "Tolerance breached (Red) or strategic impact", "SC meeting / ad hoc", "Next SC or 48 h ad hoc"]]
    for j, row in enumerate(rows):
        for i, v in enumerate(row, start=1):
            c = e.cell(5 + j, i, v); c.border = BORDER; c.font = font(size=9.5); c.alignment = Alignment(wrap_text=True, vertical="top")
            if j % 2: c.fill = FILL(B["mist"])
    bk.save(path)


# ============================================================ 14 RACI
def build_14(path):
    bk = Book("14", ["Rows = deliverables/activities (aligned with WBS), columns = roles. Rename role columns to match your project organization.",
                     "Use R (Responsible), A (Accountable), C (Consulted), I (Informed). Exactly one A per row and at least one R; the Check column validates this.",
                     "Combined 'A/R' is allowed when the accountable person also does the work."])
    roles = ["Sponsor", "Steering Committee", "Project Manager", "Business Owner", "Business Analyst", "Technical Lead",
             "Delivery Team", "QA Lead", "Procurement", "PMO", "Finance", "End Users"]
    nroles = len(roles)
    cols = [dict(h="WBS / ref.", w=9), dict(h="Deliverable / activity", w=36)]
    cols += [dict(h=r, w=9, list=("RACI", ["R", "A", "C", "I", "A/R"]), align="center") for r in roles]
    first, last = get_column_letter(3), get_column_letter(2 + nroles)
    cols += [dict(h="# A", w=6, align="center", f=f'=IF(B{{r}}="","",COUNTIF({first}{{r}}:{last}{{r}},"A")+COUNTIF({first}{{r}}:{last}{{r}},"A/R"))'),
             dict(h="# R", w=6, align="center", f=f'=IF(B{{r}}="","",COUNTIF({first}{{r}}:{last}{{r}},"R")+COUNTIF({first}{{r}}:{last}{{r}},"A/R"))'),
             dict(h="Check", w=8, align="center", f=f'=IF(B{{r}}="","",IF(AND({get_column_letter(3 + nroles)}{{r}}=1,{get_column_letter(4 + nroles)}{{r}}>=1),"OK","Fix"))')]
    #           Spon SC  PM   BO   BA   TL   Team QA  Proc PMO  Fin  Users
    acts = [
        ("1.1.1", "Project Charter", ["A", "I", "R", "C", "", "", "", "", "", "C", "C", ""]),
        ("1.1.1", "Stakeholder Register", ["I", "", "A/R", "C", "C", "", "", "", "", "", "", ""]),
        ("1.1.2", "Integrated Project Management Plan", ["A", "C", "R", "C", "C", "C", "I", "C", "C", "C", "C", ""]),
        ("1.1.2", "Scope baseline (scope statement, WBS)", ["A", "I", "R", "C", "C", "C", "C", "", "", "", "", ""]),
        ("1.1.2", "Schedule baseline", ["A", "I", "R", "C", "", "C", "C", "", "", "C", "", ""]),
        ("1.1.2", "Cost baseline & budget", ["A", "I", "R", "C", "", "C", "", "", "C", "", "C", ""]),
        ("1.1.2", "Risk Management Plan & Register", ["C", "I", "A/R", "C", "C", "C", "C", "C", "", "C", "", ""]),
        ("1.2.2", "Requirements specification & RTM", ["I", "", "C", "A", "R", "C", "C", "C", "", "", "", "C"]),
        ("1.3", "Solution design", ["I", "", "C", "C", "C", "A/R", "C", "C", "", "", "", ""]),
        ("1.4", "Procurement & contract award", ["A", "I", "R", "C", "", "C", "", "", "R", "", "C", ""]),
        ("1.4", "Build deliverables", ["", "", "A", "I", "C", "R", "R", "C", "", "", "", ""]),
        ("1.5", "Testing & quality control", ["", "", "A", "C", "C", "C", "R", "R", "", "", "", "C"]),
        ("1.5.2", "User acceptance & deliverable sign-off", ["I", "I", "R", "A", "C", "C", "", "C", "", "", "", "R"]),
        ("1.1.3", "Change requests (CCB decisions)", ["A", "C", "R", "C", "", "C", "", "", "", "I", "C", ""]),
        ("1.1.3", "Project Status Report", ["I", "I", "A/R", "I", "", "", "", "", "", "I", "", ""]),
        ("1.6", "Training & change management", ["I", "", "A", "C", "C", "", "R", "", "", "", "", "I"]),
        ("1.6.3", "Operational handover", ["C", "I", "R", "A", "", "C", "C", "", "", "", "", "I"]),
        ("1.1.4", "Lessons learned", ["I", "I", "A/R", "C", "C", "C", "C", "C", "", "C", "", ""]),
        ("1.1.4", "Project Closure Report", ["A", "I", "R", "C", "", "", "", "", "C", "C", "C", ""]),
    ]
    pre = [[w, n] + v for w, n, v in acts]
    ws = bk.register("RACI", cols, prefill=pre, n_rows=40, freeze_col=2)
    lr = ws._last_row
    rng = f"C5:{last}{lr}"
    bk.text_format(ws, rng, {"A": (B["navy"], "FFFFFF"), "A/R": (B["navy"], "FFFFFF"), "R": (B["teal"], "FFFFFF"),
                             "C": (B["amber_l"], "8A6200"), "I": (B["sky"], B["navy"])})
    chk = get_column_letter(5 + nroles)
    bk.text_format(ws, f"{chk}5:{chk}{lr}", {"OK": (B["green_l"], B["green"]), "Fix": (B["red_l"], B["red"])})
    for i in range(3, 3 + nroles):
        ws.cell(4, i).alignment = Alignment(text_rotation=90, horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 100
    # legend
    lg = bk.wb.create_sheet("Legend", 2)
    bk.band(lg, "RACI Legend", 3)
    lg.column_dimensions["A"].width = 8; lg.column_dimensions["B"].width = 18; lg.column_dimensions["C"].width = 70
    for j, (k, n, dsc, fill, fc) in enumerate([
            ("R", "Responsible", "Does the work to complete the task. At least one per row.", B["teal"], "FFFFFF"),
            ("A", "Accountable", "Ultimately answerable; approves/signs off. Exactly one per row.", B["navy"], "FFFFFF"),
            ("C", "Consulted", "Provides input before the work/decision (two-way).", B["amber_l"], "8A6200"),
            ("I", "Informed", "Kept up to date on progress or decisions (one-way).", B["sky"], B["navy"]),
            ("A/R", "Accountable & Responsible", "Same person owns and performs the work.", B["navy"], "FFFFFF")]):
        c = lg.cell(5 + j, 1, k); c.fill = FILL(fill); c.font = font(bold=True, color=fc); c.alignment = Alignment(horizontal="center"); c.border = BORDER
        lg.cell(5 + j, 2, n).font = font(bold=True); lg.cell(5 + j, 2).border = BORDER
        lg.cell(5 + j, 3, dsc).border = BORDER
    bk.save(path)


# ============================================================ 16 RISK REGISTER + 5x5
def build_16(path):
    bk = Book("16", ["Describe each risk with cause → event → effect ('Because of…, … may occur, resulting in…').",
                     "Score probability and impact 1–5 using the scales in the Risk Management Plan (doc 15). Score = P × I; rating bands: 1–4 Low, 5–9 Medium, 10–15 High, 16–25 Critical.",
                     "The 5×5 Matrix sheet counts open risks per cell (inherent and residual). Materialized risks move to the Issues log in the RAID Log (doc 19)."])
    rbs = ("RBS", ["Technical", "External", "Organizational", "Project management", "Health, safety & environment", "Commercial / contractual"])
    strat = ("Strategy", ["Avoid", "Transfer", "Mitigate", "Accept", "Escalate", "Exploit", "Share", "Enhance"])
    rstat = ("RiskStatus", ["Open", "Monitoring", "Occurred (issue)", "Closed"])
    cols = [dict(h="Risk ID", w=8), dict(h="Date raised", w=11, fmt=DATE), dict(h="Category (RBS)", w=15, list=rbs),
            dict(h="Threat / opportunity", w=11, list=("TO", ["Threat", "Opportunity"])),
            dict(h="Cause", w=24), dict(h="Risk event", w=28), dict(h="Effect on objectives", w=24), dict(h="Trigger / early warning", w=20),
            dict(h="Prob. (1-5)", w=7, list=SCORE, align="center"), dict(h="Impact (1-5)", w=7, list=SCORE, align="center"),
            dict(h="Score", w=7, align="center", f='=IF(OR(I{r}="",J{r}=""),"",I{r}*J{r})', bold=True),
            dict(h="Rating", w=10, align="center", f=rating_f("K")),
            dict(h="Proximity", w=10, list=("Prox", ["< 1 month", "1-3 months", "3-6 months", "> 6 months"])),
            dict(h="Response strategy", w=11, list=strat), dict(h="Response actions", w=30), dict(h="Risk owner", w=14),
            dict(h="Action due", w=11, fmt=DATE),
            dict(h="Residual P", w=7, list=SCORE, align="center"), dict(h="Residual I", w=7, list=SCORE, align="center"),
            dict(h="Residual score", w=8, align="center", f='=IF(OR(R{r}="",S{r}=""),"",R{r}*S{r})', bold=True),
            dict(h="Residual rating", w=10, align="center", f=rating_f("T")),
            dict(h="Contingency / fallback plan", w=24), dict(h=f"Cost impact if occurs ({CUR})", w=12, fmt=MONEY),
            dict(h="EMV (P% × cost)", w=12, fmt=MONEY, f='=IF(OR(I{r}="",W{r}=""),"",CHOOSE(I{r},0.1,0.3,0.5,0.7,0.9)*W{r}*IF(D{r}="Opportunity",-1,1))'),
            dict(h="Status", w=11, list=rstat), dict(h="Last reviewed", w=11, fmt=DATE)]
    ex = [["R-001", D, "External", "Threat", "Single supplier for key component", "Supplier delivers late",
           "Build phase delayed 3–4 weeks; cost of idle team", "Supplier misses interim milestone", 3, 4, None, None,
           "1-3 months", "Mitigate", "Weekly supplier check-ins; identify alternate supplier; milestone payments", "[PM]",
           dt.date(2026, 12, 1), 2, 3, None, None, "Engage alternate supplier; re-sequence work", 40000, None, "Open", D],
          ["R-002", D, "Organizational", "Opportunity", "Shared platform being built by another program",
           "Reuse of platform components", "Reduce build effort by ~15%", "Platform release confirmed", 2, 3, None, None,
           "3-6 months", "Enhance", "Agree reuse with program lead; align architecture", "[Tech Lead]", dt.date(2027, 1, 15),
           3, 3, None, None, "", 25000, None, "Monitoring", D]]
    ws = bk.register("Risk Register", cols, ex, n_rows=80, freeze_col=2, tab=B["red"])
    lr = ws._last_row
    bk.rating_format(ws, f"L5:L{lr}"); bk.rating_format(ws, f"U5:U{lr}")
    for col in ("K", "T"):
        rng = f"{col}5:{col}{lr}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({col}5),{col}5>=16)"], fill=FILL(B["red"]), font=Font(name="Calibri", bold=True, color="FFFFFF"), stopIfTrue=True))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({col}5),{col}5>=10)"], fill=FILL("F5B7A8"), stopIfTrue=True))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({col}5),{col}5>=5)"], fill=FILL(B["amber_l"]), stopIfTrue=True))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"ISNUMBER({col}5)"], fill=FILL(B["green_l"])))

    # 5x5 matrix sheet
    m = bk.wb.create_sheet("5x5 Matrix", 1)
    bk.band(m, "Probability × Impact Heat Map (open risks)", 15)
    m.column_dimensions["A"].width = 3
    for col in "BCDEFGHIJKLMNOP":
        m.column_dimensions[col].width = 9
    m.column_dimensions["B"].width = 12; m.column_dimensions["I"].width = 4; m.column_dimensions["J"].width = 12

    def band_fill(s):
        if s >= 16: return B["red"], "FFFFFF"
        if s >= 10: return "E67E22", "FFFFFF"
        if s >= 5: return B["amber"], B["ink"]
        return B["green"], "FFFFFF"
    RR = f"'Risk Register'!"
    for title, c0, pc, ic in (("INHERENT (before response)", 2, "I", "J"), ("RESIDUAL (after response)", 10, "R", "S")):
        cell = m.cell(4, c0, title); cell.font = font(bold=True, color=B["teal"], size=11)
        m.cell(5, c0 + 3, "IMPACT →").font = font(bold=True, color=B["navy"], size=9)
        for i in range(1, 6):
            h = m.cell(6, c0 + i, i); h.font = font(bold=True, color="FFFFFF"); h.fill = HDR_FILL
            h.alignment = Alignment(horizontal="center"); h.border = BORDER
        lab = m.cell(6, c0, "PROB. ↓"); lab.font = font(bold=True, color="FFFFFF", size=9); lab.fill = HDR_FILL
        for rr in range(5):
            p = 5 - rr
            r = 7 + rr
            m.row_dimensions[r].height = 40
            pl = m.cell(r, c0, p); pl.font = font(bold=True, color="FFFFFF"); pl.fill = HDR_FILL
            pl.alignment = Alignment(horizontal="center", vertical="center"); pl.border = BORDER
            for i in range(1, 6):
                s = p * i
                fill, fc = band_fill(s)
                f = (f'=COUNTIFS({RR}${pc}$5:${pc}${lr},{p},{RR}${ic}$5:${ic}${lr},{i},{RR}$Y$5:$Y${lr},"<>Closed")')
                c = m.cell(r, c0 + i, f)
                c.fill = FILL(fill); c.font = font(bold=True, size=16, color=fc); c.border = BORDER
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.number_format = '0;-0;""'
    # rating summary
    r = 14
    m.cell(r, 2, "Open risks by rating").font = font(bold=True, color=B["teal"], size=11)
    for i, h in enumerate(["Rating", "Inherent", "Residual", "Action required"]):
        c = m.cell(r + 1, 2 + i * (1 if i < 3 else 1), h)
    heads = [("B", "Rating"), ("C", "Inherent"), ("D", "Residual"), ("E", "Action required")]
    for col, h in heads:
        c = m[f"{col}{r + 1}"]; c.value = h; c.font = font(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.border = BORDER
    m.merge_cells(start_row=r + 1, start_column=5, end_row=r + 1, end_column=9)
    acts = {"Critical": "Escalate to Sponsor/SC within 24 h; weekly review", "High": "Response plan within 5 days; weekly review",
            "Medium": "Response plan; monthly review", "Low": "Accept & monitor (watch list)"}
    for j, rt in enumerate(["Critical", "High", "Medium", "Low"]):
        rr = r + 2 + j
        c = m.cell(rr, 2, rt); c.border = BORDER; c.font = font(bold=True)
        m.cell(rr, 3, f'=COUNTIFS({RR}$L$5:$L${lr},"{rt}",{RR}$Y$5:$Y${lr},"<>Closed")').border = BORDER
        m.cell(rr, 4, f'=COUNTIFS({RR}$U$5:$U${lr},"{rt}",{RR}$Y$5:$Y${lr},"<>Closed")').border = BORDER
        m.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=9)
        m.cell(rr, 5, acts[rt]).font = font(size=9)
    bk.rating_format(m, f"B{r + 2}:B{r + 5}")
    m.cell(r + 7, 2, "Total EMV of open risks (contingency basis):").font = font(bold=True)
    c = m.cell(r + 7, 7, f'=SUMIFS({RR}$X$5:$X${lr},{RR}$Y$5:$Y${lr},"<>Closed")'); c.number_format = MONEY
    c.font = font(bold=True, size=12, color=B["navy"]); c.fill = FILL("EAF2F2")

    # scales sheet
    sc = bk.wb.create_sheet("Scales", 2)
    bk.band(sc, "Probability & Impact Scales (from doc 15)", 6)
    for col, w in zip("ABCDEF", (12, 14, 14, 18, 22, 26)):
        sc.column_dimensions[col].width = w
    heads = ["Score", "Probability", "Cost impact", "Schedule impact", "Scope / quality", "Safety / reputation"]
    data = [[5, "> 80%", "> 20% BAC", "> 20% / misses gate", "Product unusable", "Serious injury / national media"],
            [4, "61–80%", "10–20%", "10–20%", "Unacceptable to sponsor", "Lost-time injury / regulator"],
            [3, "41–60%", "5–10%", "5–10%", "Major areas affected", "Medical aid / local media"],
            [2, "21–40%", "1–5%", "1–5%", "Minor areas affected", "First aid / internal complaint"],
            [1, "≤ 20%", "< 1%", "< 1%", "Barely noticeable", "Negligible"]]
    for i, h in enumerate(heads, start=1):
        c = sc.cell(4, i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.border = BORDER
    for j, row in enumerate(data):
        for i, v in enumerate(row, start=1):
            c = sc.cell(5 + j, i, v); c.border = BORDER; c.font = font(size=9.5); c.alignment = Alignment(wrap_text=True)
    sc.cell(11, 1, "EMV uses probability mid-points: 1=10%, 2=30%, 3=50%, 4=70%, 5=90%. Opportunities produce negative EMV (reduce contingency).").font = font(size=9, italic=True, color=B["grey"])
    bk.save(path)


# ============================================================ 19 RAID LOG
def build_19(path):
    bk = Book("19", ["Four logs in one workbook: Risks (summary view — the full register is doc 16), Assumptions, Issues and Dependencies.",
                     "Review the RAID Log in the weekly team meeting; the Dashboard counts open items by priority and age.",
                     "Escalate issues per the escalation path in the Integrated PMP (doc 02) and Communication Plan (doc 12)."])
    rcols = [dict(h="Risk ID (doc 16)", w=10), dict(h="Risk description", w=40), dict(h="Score", w=7, list=("Score25", [1, 2, 3, 4, 5, 6, 8, 9, 10, 12, 15, 16, 20, 25]), align="center"),
             dict(h="Rating", w=10, f=rating_f("C"), align="center"), dict(h="Owner", w=14), dict(h="Response summary", w=36),
             dict(h="Status", w=11, list=OPEN_STATUS), dict(h="Last updated", w=11, fmt=DATE)]
    ws = bk.register("Risks", rcols, [["R-001", "Supplier delivers key component late", 12, None, "[PM]", "Weekly check-ins; alternate supplier", "Open", D]], n_rows=50, tab=B["red"])
    bk.rating_format(ws, f"D5:D{ws._last_row}"); status_fmt(bk, ws, f"G5:G{ws._last_row}")
    acols = [dict(h="Assumption ID", w=10), dict(h="Assumption", w=40), dict(h="Category", w=14, list=("ACat", ["Resource", "Technical", "Financial", "Schedule", "External", "Business"])),
             dict(h="Source", w=14), dict(h="Impact if false", w=30), dict(h="Owner", w=14), dict(h="Validate by", w=11, fmt=DATE),
             dict(h="Validation status", w=12, list=("AVal", ["Not validated", "Validated", "Invalid – raise risk/issue"])),
             dict(h="Status", w=10, list=OPEN_STATUS)]
    ws = bk.register("Assumptions", acols, [["A-001", "Business SMEs available 4 h/week during requirements phase", "Resource", "Charter",
                                            "Requirements delayed; schedule slip", "[Business Owner]", dt.date(2026, 11, 20), "Not validated", "Open"]], n_rows=50, tab=B["amber"])
    status_fmt(bk, ws, f"I5:I{ws._last_row}")
    icols = [dict(h="Issue ID", w=8), dict(h="Date raised", w=11, fmt=DATE), dict(h="Issue description", w=36),
             dict(h="Category", w=14, list=("ICat", ["Scope", "Schedule", "Cost", "Quality", "Resource", "Stakeholder", "Technical", "Vendor"])),
             dict(h="Priority", w=10, list=("Priority", PRIORITY)), dict(h="Impact", w=28), dict(h="Raised by", w=13), dict(h="Owner", w=13),
             dict(h="Action / resolution", w=32), dict(h="Target date", w=11, fmt=DATE),
             dict(h="Escalation level", w=11, list=("Esc", ["1 - PM", "2 - Sponsor", "3 - Steering Committee"])),
             dict(h="Status", w=10, list=OPEN_STATUS), dict(h="Date closed", w=11, fmt=DATE),
             dict(h="Age (days)", w=8, align="center", f='=IF(B{r}="","",IF(M{r}<>"",M{r}-B{r},TODAY()-B{r}))'),
             dict(h="Overdue", w=8, align="center", f='=IF(OR(J{r}="",L{r}="Closed"),"",IF(TODAY()>J{r},"Yes",""))'),
             dict(h="Related risk / CR", w=12)]
    ws = bk.register("Issues", icols, [["I-001", D, "Test environment not available as planned", "Technical", "High",
                                       "Testing start at risk (5 days)", "[QA Lead]", "[Tech Lead]", "Escalate to IT ops; use temporary environment",
                                       dt.date(2026, 11, 17), "1 - PM", "Open", None, None, None, "R-003"]], n_rows=60, tab=B["navy"])
    lr = ws._last_row
    bk.text_format(ws, f"E5:E{lr}", {"Critical": (B["red"], "FFFFFF"), "High": ("F5B7A8", "7A1E10"), "Medium": (B["amber_l"], "8A6200"), "Low": (B["green_l"], B["green"])})
    status_fmt(bk, ws, f"L5:L{lr}")
    bk.text_format(ws, f"O5:O{lr}", {"Yes": (B["red_l"], B["red"])})
    dcols = [dict(h="Dependency ID", w=10), dict(h="Dependency description", w=36),
             dict(h="Type", w=14, list=("DType", ["Internal", "External", "Mandatory", "Discretionary"])),
             dict(h="Direction", w=12, list=("Dir", ["Inbound (we need)", "Outbound (others need)"])),
             dict(h="Other party / project", w=18), dict(h="Deliverable / WBS affected", w=16), dict(h="Needed by", w=11, fmt=DATE),
             dict(h="Owner", w=13), dict(h="Health (RAG)", w=9, list=("RAG", ["Green", "Amber", "Red"])), dict(h="Status", w=10, list=OPEN_STATUS),
             dict(h="Comments", w=26)]
    ws = bk.register("Dependencies", dcols, [["D-001", "Network upgrade by IT Infrastructure project", "External", "Inbound (we need)",
                                             "[Infra project]", "1.6.2", dt.date(2027, 4, 1), "[PM]", "Green", "Open", ""]], n_rows=40, tab=B["teal"])
    bk.rag_format(ws, f"I5:I{ws._last_row}"); status_fmt(bk, ws, f"J5:J{ws._last_row}")
    # dashboard
    db = bk.wb.create_sheet("Dashboard", 1)
    bk.band(db, "RAID Dashboard", 6)
    db.column_dimensions["A"].width = 3
    for col, w in zip("BCDEFG", (26, 10, 10, 10, 10, 10)):
        db.column_dimensions[col].width = w
    heads = ["Log", "Open", "In progress", "On hold", "Closed", "Total"]
    for i, h in enumerate(heads):
        c = db.cell(4, 2 + i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.border = BORDER
    logs = [("Risks", "G"), ("Assumptions", "I"), ("Issues", "L"), ("Dependencies", "J")]
    for j, (name, col) in enumerate(logs):
        r = 5 + j
        bk.label(db, r, 2, name, bold=True)
        for i, st in enumerate(["Open", "In progress", "On hold", "Closed"]):
            bk.value(db, r, 3 + i, f'=COUNTIF({name}!${col}$5:${col}$200,"{st}")', fx=True)
        bk.value(db, r, 7, f'=COUNTA({name}!$A$5:$A$200)', fx=True, bold=True)
    r = 11
    bk.label(db, r, 2, "Open issues by priority", color=B["teal"], size=11).border = None
    for j, p in enumerate(PRIORITY):
        bk.label(db, r + 1 + j, 2, p, bold=False, color=B["ink"])
        bk.value(db, r + 1 + j, 3, f'=COUNTIFS(Issues!$E$5:$E$200,"{p}",Issues!$L$5:$L$200,"<>Closed")', fx=True)
    bk.label(db, r + 6, 2, "Overdue issues", bold=True); bk.value(db, r + 6, 3, '=COUNTIF(Issues!$O$5:$O$200,"Yes")', fx=True, bold=True)
    bk.label(db, r + 7, 2, "Average age of open issues (d)", bold=True)
    bk.value(db, r + 7, 3, '=IFERROR(AVERAGEIFS(Issues!$N$5:$N$200,Issues!$L$5:$L$200,"<>Closed",Issues!$A$5:$A$200,"<>"),0)', fmt="0.0", fx=True)
    bk.label(db, r + 8, 2, "Red dependencies", bold=True); bk.value(db, r + 8, 3, '=COUNTIF(Dependencies!$I$5:$I$200,"Red")', fx=True, bold=True)
    ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100; ch.title = "RAID items by status"; ch.style = 10
    ch.add_data(Reference(db, min_col=3, max_col=6, min_row=4, max_row=8), titles_from_data=True)
    ch.set_categories(Reference(db, min_col=2, min_row=5, max_row=8))
    ch.height, ch.width = 7.5, 14
    db.add_chart(ch, "I4")
    bk.save(path)


# ============================================================ 21 CHANGE LOG
def build_21(path):
    bk = Book("21", ["Log every change request (doc 20) on submission and assign the next CR ID.",
                     "Record quantified impacts; approved changes update baselines and trigger a major version of affected documents.",
                     "The Summary sheet totals approved schedule and cost impacts against the original baselines."])
    st = ("CRStatus", ["Submitted", "Under analysis", "Pending decision", "Approved", "Approved with conditions", "Rejected", "Deferred", "Implemented", "Closed"])
    cols = [dict(h="CR ID", w=10), dict(h="Title", w=28), dict(h="Requested by", w=14), dict(h="Date submitted", w=11, fmt=DATE),
            dict(h="Category", w=12, list=("CRCat", ["Scope", "Schedule", "Cost", "Quality", "Resource", "Contract", "Document"])),
            dict(h="Description / reason", w=36), dict(h="Priority", w=10, list=("Priority", PRIORITY)),
            dict(h="Schedule impact (wd)", w=10, fmt="0;-0;0"), dict(h=f"Cost impact ({CUR})", w=12, fmt=MONEY),
            dict(h="Baselines affected", w=14, list=("Base", ["None", "Scope", "Schedule", "Cost", "Scope+Schedule", "Schedule+Cost", "Scope+Schedule+Cost"])),
            dict(h="Decision authority", w=14, list=("Auth", ["Project Manager", "CCB", "Steering Committee"])),
            dict(h="Status", w=14, list=st), dict(h="Decision date", w=11, fmt=DATE), dict(h="Decision ref. (doc 23)", w=11),
            dict(h="Implemented date", w=11, fmt=DATE), dict(h="Documents re-versioned", w=22),
            dict(h="Days to decision", w=9, align="center", f='=IF(OR(D{r}="",M{r}=""),"",M{r}-D{r})')]
    ex = [["CR-001", "Add mobile access to reporting module", "[Business Owner]", D, "Scope", "Users require field access", "Medium",
           8, 18000, "Scope+Schedule+Cost", "CCB", "Pending decision", None, "", None, ""]]
    ws = bk.register("Change Log", cols, ex, n_rows=60, freeze_col=2)
    lr = ws._last_row
    bk.text_format(ws, f"L5:L{lr}", {"Approved": (B["green_l"], B["green"]), "Approved with conditions": (B["green_l"], B["green"]),
                                      "Rejected": (B["red_l"], B["red"]), "Pending decision": (B["amber_l"], "8A6200"),
                                      "Implemented": (B["sky"], B["navy"]), "Closed": ("E5E8EB", B["grey"])})
    s = bk.wb.create_sheet("Summary", 1)
    bk.band(s, "Change Summary", 4)
    s.column_dimensions["A"].width = 3; s.column_dimensions["B"].width = 30; s.column_dimensions["C"].width = 14
    s.column_dimensions["D"].width = 4; s.column_dimensions["E"].width = 34; s.column_dimensions["F"].width = 16
    summary_block(bk, s, 4, 2, "Change requests by status",
                  [(v, f'=COUNTIF(\'Change Log\'!$L$5:$L${lr},"{v}")') for v in st[1]] + [("Total", f"=COUNTA('Change Log'!$A$5:$A${lr})")])
    CL = "'Change Log'!"
    appr = lambda col: f'=SUMIFS({CL}${col}$5:${col}${lr},{CL}$L$5:$L${lr},"Approved*")+SUMIFS({CL}${col}$5:${col}${lr},{CL}$L$5:$L${lr},"Implemented")+SUMIFS({CL}${col}$5:${col}${lr},{CL}$L$5:$L${lr},"Closed")'
    summary_block(bk, s, 4, 5, "Approved impact on baselines", [
        ("Approved schedule impact (wd)", appr("H"), "0"),
        (f"Approved cost impact ({CUR})", appr("I"), MONEY),
        (f"Pending cost exposure ({CUR})", f'=SUMIFS({CL}$I$5:$I${lr},{CL}$L$5:$L${lr},"Pending decision")+SUMIFS({CL}$I$5:$I${lr},{CL}$L$5:$L${lr},"Under analysis")', MONEY),
        ("Average days to decision", f"=IFERROR(AVERAGE({CL}$Q$5:$Q${lr}),0)", "0.0"),
    ], header=("Metric", "Value"))
    bk.save(path)


# ============================================================ 23 DECISION LOG
def build_23(path):
    bk = Book("23", ["Record significant project decisions with context, options and rationale so they are traceable and not reopened.",
                     "Reference related change requests (doc 21), risks/issues (docs 16/19) and the authority level that made the decision.",
                     "Communicate each decision per the Communication Plan (doc 12)."])
    cols = [dict(h="Decision ID", w=10), dict(h="Date", w=11, fmt=DATE), dict(h="Decision title", w=26), dict(h="Context / problem", w=32),
            dict(h="Options considered", w=30), dict(h="Decision made", w=30), dict(h="Rationale", w=30),
            dict(h="Decision maker(s)", w=16), dict(h="Authority level", w=14, list=("Auth", ["Project Manager", "CCB", "Sponsor", "Steering Committee"])),
            dict(h="Performance domain", w=13, list=("Domain", DOMAINS)),
            dict(h="Impact (scope / schedule / cost / risk)", w=24), dict(h="Related CR / risk / issue", w=12),
            dict(h="Communicated to", w=18), dict(h="Status", w=11, list=("DStatus", ["Proposed", "Approved", "Implemented", "Superseded", "Revoked"])),
            dict(h="Review date", w=11, fmt=DATE)]
    ex = [["DEC-001", D, "Adopt phased deployment", "Big-bang go-live risk too high for operations",
           "A) Big-bang  B) Phased by region  C) Pilot + big-bang", "Option B – phased by region", "Lower operational risk; manageable training load",
           "Steering Committee", "Steering Committee", "Risk", "Schedule +10 wd; risk score R-004 reduced", "CR-002 / R-004", "All stakeholders (C-08)", "Approved", dt.date(2027, 2, 1)]]
    ws = bk.register("Decision Log", cols, ex, n_rows=60, freeze_col=3)
    bk.save(path)


# ============================================================ 24 LESSONS LEARNED
def build_24(path):
    bk = Book("24", ["Capture lessons continuously (not only at closure) — both what worked well and what should be improved.",
                     "Tag each lesson by focus area and PMBOK 8 performance domain so it can be reused by future projects and the PMO.",
                     "The Summary sheet aggregates lessons by domain; top lessons feed the Project Closure Report (doc 25)."])
    cols = [dict(h="LL ID", w=8), dict(h="Date", w=11, fmt=DATE), dict(h="Focus area / phase", w=16, list=("Focus", FOCUS_AREAS)),
            dict(h="Performance domain", w=13, list=("Domain", DOMAINS)),
            dict(h="Type", w=12, list=("LLType", ["Went well", "To improve"])),
            dict(h="Situation / what happened", w=34), dict(h="Impact on project", w=26), dict(h="Root cause", w=26),
            dict(h="Recommendation for future projects", w=34), dict(h="Action owner", w=14),
            dict(h="Action status", w=11, list=("LLAct", ["Proposed", "In progress", "Implemented", "Not actioned"])),
            dict(h="Applicability", w=14, list=("Appl", ["This project", "Program", "Organization (PMO)"])),
            dict(h="Submitted by", w=14), dict(h="Keywords", w=16)]
    ex = [["LL-001", D, "Planning", "Stakeholders", "To improve", "End-user group engaged only at UAT",
           "Late change requests; 2-week UAT extension", "Stakeholder analysis missed operational users",
           "Include end-user representatives from requirements phase; use power/interest grid at Gate 1", "[PM]", "Proposed",
           "Organization (PMO)", "[Name]", "engagement, UAT"]]
    ws = bk.register("Lessons Learned", cols, ex, n_rows=60, freeze_col=2)
    lr = ws._last_row
    bk.text_format(ws, f"E5:E{lr}", {"Went well": (B["green_l"], B["green"]), "To improve": (B["amber_l"], "8A6200")})
    s = bk.wb.create_sheet("Summary", 1)
    bk.band(s, "Lessons by Performance Domain", 5)
    for col, w in zip("ABCDE", (3, 22, 12, 12, 10)):
        s.column_dimensions[col].width = w
    heads = ["Performance domain", "Went well", "To improve", "Total"]
    for i, h in enumerate(heads):
        c = s.cell(4, 2 + i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.border = BORDER
    LL = "'Lessons Learned'!"
    for j, dm in enumerate(DOMAINS):
        r = 5 + j
        bk.label(s, r, 2, dm, bold=False, color=B["ink"])
        bk.value(s, r, 3, f'=COUNTIFS({LL}$D$5:$D${lr},$B{r},{LL}$E$5:$E${lr},"Went well")', fx=True)
        bk.value(s, r, 4, f'=COUNTIFS({LL}$D$5:$D${lr},$B{r},{LL}$E$5:$E${lr},"To improve")', fx=True)
        bk.value(s, r, 5, f"=C{r}+D{r}", fx=True, bold=True)
    ch = BarChart(); ch.type = "bar"; ch.grouping = "stacked"; ch.overlap = 100; ch.title = "Lessons by domain"; ch.style = 10
    ch.add_data(Reference(s, min_col=3, max_col=4, min_row=4, max_row=4 + len(DOMAINS)), titles_from_data=True)
    ch.set_categories(Reference(s, min_col=2, min_row=5, max_row=4 + len(DOMAINS)))
    ch.height, ch.width = 7.5, 14
    s.add_chart(ch, "G4")
    bk.save(path)
