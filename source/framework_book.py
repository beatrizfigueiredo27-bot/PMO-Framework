"""00 - Master PMO Documentation Framework.xlsx"""
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

import config
from config import DOMAINS, FOCUS_AREAS, GLOSSARY, VERSIONING
from framework import artefacts, GATES, CATS, AUTHORITY, ESCALATION, RAG, DOC_STANDARD
from xlsx_kit import Book, FILL, font, BORDER, B, HDR_FILL, FX_FILL


def simple_table(bk, ws, r0, headers, rows, widths=None, wrap=True):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(r0, i, h); c.font = font(bold=True, color="FFFFFF", size=9.5); c.fill = HDR_FILL; c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    for j, row in enumerate(rows):
        for i, v in enumerate(row, start=1):
            c = ws.cell(r0 + 1 + j, i, v); c.border = BORDER; c.font = font(size=9.5)
            c.alignment = Alignment(wrap_text=wrap, vertical="top")
            if j % 2:
                c.fill = FILL(B["mist"])
    return r0 + 1 + len(rows)


def build_framework(path):
    entry = ("00", "Master PMO Documentation Framework", "MDF", "xlsx", list(DOMAINS), list(FOCUS_AREAS), "PMO Lead")
    config.REGISTER.insert(0, entry)
    try:
        bk = Book("00", ["The Document Register defines, for all 25 artefacts, whether each is Mandatory (M), Conditional (C) or Optional (O) by project complexity, when it is created, who owns and approves it, when it is baselined and which artefacts it feeds.",
                         "Use the Gate Checklist at each governance gate: an artefact marked M for your complexity level must be at the stated status before the gate can pass.",
                         "Delegated authority, escalation and RAG thresholds here are the PMO defaults; record project-specific tailoring in the Integrated PMP (doc 02)."])
    finally:
        config.REGISTER.remove(entry)
    arts = artefacts()
    cls = ("Cls", ["M", "C", "O"])
    cols = [dict(h="#", w=5, align="center"), dict(h="Artefact", w=30), dict(h="Document ID", w=17), dict(h="Format", w=7, align="center"),
            dict(h="Primary performance domain", w=14), dict(h="Focus area(s)", w=20),
            dict(h="Low", w=6, list=cls, align="center"), dict(h="Medium", w=7, list=cls, align="center"), dict(h="High", w=6, list=cls, align="center"),
            dict(h="Condition / tailoring note", w=34), dict(h="Control category", w=18), dict(h="Created in", w=14), dict(h="Trigger", w=24),
            dict(h="Owner", w=16), dict(h="Approver", w=20), dict(h="Consulted", w=20), dict(h="Baselined", w=26),
            dict(h="Update cadence", w=24), dict(h="Feeds (outputs to)", w=18), dict(h="Fed by (inputs from)", w=18),
            dict(h="Objective acceptance criteria", w=60)]
    pre = []
    for a in arts:
        pre.append([a["no"], a["title"], a["id"], a["ext"], a["domains"][0], ", ".join(a["focus"]),
                    a["tailor"]["Low"], a["tailor"]["Medium"], a["tailor"]["High"], a["cond"], a["catName"],
                    f'{a["created"]} ({a["gate"]})', a["trigger"], a["owner"], a["approver"], a["consulted"],
                    a["baseline"], a["cadence"], ", ".join(a["feeds"]) or "Organizational process assets",
                    ", ".join(a["fedBy"]) or "Business case / project request",
                    "\n".join("• " + x for x in a["accept"])])
    ws = bk.register("Document Register", cols, prefill=pre, n_rows=len(pre), freeze_col=2)
    for r in range(5, 5 + len(pre)):
        ws.row_dimensions[r].height = 92
    bk.text_format(ws, f"G5:I{4 + len(pre)}", {"M": (B["navy"], "FFFFFF"), "C": (B["amber_l"], "8A6200"), "O": ("E5E8EB", B["grey"])})

    # Gate checklist
    g = bk.wb.create_sheet("Gate Checklist")
    bk.band(g, "Governance Gate Checklist", 9)
    r = 4
    for gid, name, fa, appr, crit in GATES:
        c = g.cell(r, 1, f"{gid} — {name}"); c.font = font(bold=True, color=B["teal"], size=12)
        g.cell(r + 1, 1, f"Focus area: {fa}   •   Approver: {appr}").font = font(size=9.5, color=B["grey"])
        g.merge_cells(start_row=r + 2, start_column=1, end_row=r + 2, end_column=9)
        cc = g.cell(r + 2, 1, "Exit criteria: " + crit); cc.font = font(size=9.5); cc.alignment = Alignment(wrap_text=True)
        g.row_dimensions[r + 2].height = 30
        rows = []
        for a in arts:
            if gate_needs(a, gid):
                rows.append([a["no"], a["title"], a["tailor"]["Low"], a["tailor"]["Medium"], a["tailor"]["High"],
                             gate_status(a, gid), a["approver"], "☐", ""])
        r = simple_table(bk, g, r + 3, ["#", "Artefact", "Low", "Med", "High", "Required status at gate", "Approver", "Done", "Evidence / link"],
                         rows, widths=[5, 34, 6, 6, 6, 30, 26, 7, 28]) + 2
    g.freeze_panes = "A4"

    # Information flow matrix
    m = bk.wb.create_sheet("Information Flow")
    bk.band(m, "Information Flow Matrix (row feeds column)", len(arts) + 2)
    m.column_dimensions["A"].width = 6; m.column_dimensions["B"].width = 30
    for j, a in enumerate(arts):
        L = get_column_letter(3 + j)
        m.column_dimensions[L].width = 4.2
        c = m.cell(4, 3 + j, a["no"]); c.font = font(bold=True, color="FFFFFF", size=9); c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center"); c.border = BORDER
    m.cell(4, 1, "From").font = font(bold=True, color="FFFFFF"); m.cell(4, 1).fill = HDR_FILL
    m.cell(4, 2, "Artefact ↓ feeds →").font = font(bold=True, color="FFFFFF"); m.cell(4, 2).fill = HDR_FILL
    for i, a in enumerate(arts):
        r = 5 + i
        m.cell(r, 1, a["no"]).font = font(bold=True); m.cell(r, 2, a["title"]).font = font(size=9.5)
        for j, b in enumerate(arts):
            c = m.cell(r, 3 + j, "●" if b["no"] in a["feeds"] else ("·" if a["no"] == b["no"] else ""))
            c.alignment = Alignment(horizontal="center"); c.border = BORDER
            c.font = font(color=B["teal"], bold=True)
            if a["no"] == b["no"]:
                c.fill = FILL(B["sky"])
    m.freeze_panes = "C5"

    # Baseline & change control
    bc = bk.wb.create_sheet("Baseline & Change Control")
    bk.band(bc, "Baseline & Change-Control Relationships", 4)
    rows = [[CATS[k][0], CATS[k][1], ", ".join(a["no"] for a in arts if a["cat"] == k)] for k in CATS]
    r = simple_table(bk, bc, 4, ["Control category", "Versioning & change rule", "Artefacts"], rows, widths=[26, 80, 30]) + 2
    bc.cell(r, 1, "Performance measurement baseline (PMB)").font = font(bold=True, color=B["teal"], size=11)
    r = simple_table(bk, bc, r + 1, ["Baseline", "Composition", "Artefacts"],
                     [["Scope baseline", "Approved scope statement + WBS + WBS Dictionary", "03, 05"],
                      ["Schedule baseline", "Approved start/finish dates of activities and milestones", "07"],
                      ["Cost baseline", "Time-phased budget incl. contingency, excl. management reserve", "09"],
                      ["PMB", "Integrated scope-schedule-cost baseline used for EVM", "05, 07, 09"]]) + 2
    bc.cell(r, 1, "Versioning scheme").font = font(bold=True, color=B["teal"], size=11)
    simple_table(bk, bc, r + 1, ["Version", "Meaning", "Rule"], [list(v) for v in VERSIONING])

    # Authority, escalation, RAG
    au = bk.wb.create_sheet("Authority & Escalation")
    bk.band(au, "Delegated Authority, Escalation & RAG Thresholds", 4)
    r = simple_table(bk, au, 4, ["Decision", "Project Manager", "Change Control Board", "Sponsor / Steering Committee"],
                     [list(x) for x in AUTHORITY], widths=[26, 34, 34, 44]) + 2
    au.cell(r, 1, "Escalation levels").font = font(bold=True, color=B["teal"], size=11)
    r = simple_table(bk, au, r + 1, ["Level", "Escalate to", "Trigger", "Response time"], [list(x) for x in ESCALATION]) + 2
    au.cell(r, 1, "RAG thresholds").font = font(bold=True, color=B["teal"], size=11)
    r2 = simple_table(bk, au, r + 1, ["Dimension", "Green", "Amber", "Red"], [list(x) for x in RAG])
    for rr in range(r + 2, r2):
        au.cell(rr, 2).fill = FILL(B["green_l"]); au.cell(rr, 3).fill = FILL(B["amber_l"]); au.cell(rr, 4).fill = FILL(B["red_l"])

    # Document standard + glossary
    st = bk.wb.create_sheet("Standards & Glossary")
    bk.band(st, "Document-Control Standard & Terminology", 2)
    r = simple_table(bk, st, 4, ["Element", "Standard"], [list(x) for x in DOC_STANDARD], widths=[26, 100]) + 2
    st.cell(r, 1, "Predictive governance terminology").font = font(bold=True, color=B["teal"], size=11)
    simple_table(bk, st, r + 1, ["Term", "Definition"], [list(x) for x in GLOSSARY])
    bk.save(path)


def gate_needs(a, gid):
    order = ["G1", "G2", "G3", "G4", "G5"]
    if a["no"] == "25":
        return gid == "G5"
    if a["cat"] in ("FORM",):
        return False
    return order.index(a["gate"]) <= order.index(gid)


def gate_status(a, gid):
    if a["gate"] == gid:
        return "Approved v1.0" if a["cat"] in ("AUTH", "BASE", "PLAN", "CLOSE") else "Established and current"
    if a["cat"] in ("BASE", "PLAN", "AUTH"):
        return "Current approved version; changes via CR"
    if a["cat"] == "RPT":
        return "Latest period issued"
    return "Up to date; gate snapshot taken" if gid != "G5" else "Final issue, archived"
